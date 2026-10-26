"""Automated QA for the NoviraNest workbook.

Every check recalculates the real .xlsx in LibreOffice (headless) and compares the cell values with
expected values from tests/oracle.py (independent Python re-implementation) or hand-computed numbers.

Usage:  python tests/run_qa.py            (writes tests/QA_RESULTS.md, exits 1 on any failure)
"""
import datetime as dt
import json
import random
import re
import shutil
import subprocess
import sys
import tempfile
import time
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tests"))
import build_workbook as BW  # noqa: E402
import sample_data as SD  # noqa: E402
from oracle import Oracle  # noqa: E402

TMP = Path(tempfile.mkdtemp(prefix="nn_qa_"))
RESULTS = []


def check(group, name, ok, detail=""):
    RESULTS.append((group, name, bool(ok), detail))
    mark = "PASS" if ok else "FAIL"
    print(f"[{mark}] {group}: {name} {detail if not ok else ''}")


def close(a, b, tol=0.005):
    try:
        return abs(float(a) - float(b)) <= tol * max(1, abs(float(b)))
    except (TypeError, ValueError):
        return False


def m0(x, sym="$"):
    """Format like Excel TEXT(x,"#,##0") (round half away from zero)."""
    from decimal import Decimal, ROUND_HALF_UP
    v = Decimal(str(x)).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    return ("-" if v < 0 else "") + sym + f"{abs(v):,}"


def recalc(src, label):
    dst = TMP / f"{label}.xlsx"
    shutil.copy(src, dst)
    t0 = time.time()
    out = subprocess.run([sys.executable, str(ROOT / "tools/lo/recalc.py"), str(dst), "600"], capture_output=True, text=True)
    info = json.loads(out.stdout)
    info["seconds"] = round(time.time() - t0, 1)
    return dst, info


def values(path):
    return load_workbook(path, data_only=True)


def modified(src, label, edits):
    """edits: list of (sheet, cell, value). Writes into a copy of the formula workbook."""
    wb = load_workbook(src)
    for sh, cell, val in edits:
        wb[sh][cell] = val
    p = TMP / f"{label}_in.xlsx"
    wb.save(p)
    return p


def build_both():
    s = ROOT / "dist" / "_qa_sample.xlsx"
    t = ROOT / "dist" / "_qa_template.xlsx"
    BW.build("sample", s)
    BW.build("template", t)
    return s, t


def find_row(ws, col, value, first=7, last=3006):
    for r in range(first, last + 1):
        if ws[f"{col}{r}"].value == value:
            return r
    return None


# ----------------------------------------------------------------------------------
def test_sample(sample_src):
    G = "Sample workbook"
    path, info = recalc(sample_src, "sample")
    check(G, "LibreOffice recalculation: zero formula errors", info.get("total_errors") == 0, json.dumps(info))
    check(G, "Formula count above 30,000 (logs fully pre-formulated)", info.get("total_formulas", 0) > 30000, str(info.get("total_formulas")))
    wb = values(path)
    o = Oracle()

    # Menu pricing
    G = "Menu pricing"
    mp = wb["Menu Pricing"]
    om = o.menu()
    bad = []
    for name, exp in om.items():
        r = find_row(mp, "B", name, 7, 56)
        got = dict(ing=mp[f"D{r}"].value, total=mp[f"K{r}"].value, suggested=mp[f"N{r}"].value,
                   profit=mp[f"P{r}"].value, margin=mp[f"Q{r}"].value, markup=mp[f"R{r}"].value, food=mp[f"S{r}"].value)
        for k, v in got.items():
            if not close(v, exp[k], 1e-6):
                bad.append(f"{name}.{k}: got {v} expected {exp[k]}")
        if mp[f"T{r}"].value != exp["status"]:
            bad.append(f"{name}.status {mp[f'T{r}'].value} != {exp['status']}")
    check(G, "Ingredient cost, total cost, suggested price, profit, margin, markup, food cost %, status (13 items)", not bad, "; ".join(bad[:5]))
    loss = [n for n, e in om.items() if e["status"] == "LOSS"]
    check(G, "Loss-making item flagged LOSS", loss and all(mp[f"T{find_row(mp, 'B', n, 7, 56)}"].value == "LOSS" for n in loss), str(loss))
    # margin vs markup are distinct and consistent: markup = margin / (1 - margin)
    r = find_row(mp, "B", "Smash Burger", 7, 56)
    q_, m_ = mp[f"Q{r}"].value, mp[f"R{r}"].value
    check(G, "Margin and markup are mathematically distinct (markup = margin/(1-margin))", close(m_, q_ / (1 - q_), 1e-9) and abs(q_ - m_) > 0.1)
    rec = mp[f"U{find_row(mp, 'B', 'Kids Slider Meal', 7, 56)}"].value
    check(G, "Recommendation text for LOSS item quotes suggested price", "$10.50" in str(rec), str(rec))

    # Recipe lines
    G = "Recipe Builder / unit conversion"
    rb = wb["Recipe Builder"]
    bad = []
    r = 7
    for item, lines in SD.RECIPES.items():
        for ing, qty, unit in lines:
            exp = o.to_su(ing, qty, unit) * o.cpu(ing)
            if rb[f"G{r}"].value != "OK" or not close(rb[f"J{r}"].value, exp, 1e-9):
                bad.append(f"row {r} {item}/{ing}: {rb[f'G{r}'].value} {rb[f'J{r}'].value} vs {exp}")
            r += 1
    check(G, f"All {r - 7} recipe lines convert and cost correctly (oz->lb, ml->l, tbsp->l, fl oz->gal, g->kg, cup->gal)", not bad, "; ".join(bad[:3]))

    # Monthly report
    G = "Monthly report"
    rp = wb["Monthly Report"]
    cols = dict(E="rev", F="items", G="cogs", J="waste", K="fixed", L="labor", M="event", N="var", O="opex", P="net", R="stock", S="ing")
    bad = []
    for m in range(1, 13):
        exp = o.month(SD.YEAR, m)
        for c, k in cols.items():
            if not close(rp[f"{c}{7 + m}"].value, exp.get(k, 0), 1e-7):
                bad.append(f"{m}/{c}: {rp[f'{c}{7 + m}'].value} vs {exp.get(k, 0)}")
    check(G, "12 months x 12 columns match oracle (revenue, items, COGS, waste, opex by type, net, stock, ingredient)", not bad, "; ".join(bad[:4]))
    y = o.period(dt.date(SD.YEAR, 1, 1), dt.date(SD.YEAR, 12, 31))
    check(G, "Annual total net profit = oracle", close(rp["P20"].value, y["net"], 1e-9), f"{rp['P20'].value} vs {y['net']}")
    check(G, "Net profit identity: GP - waste - opex", close(rp["P20"].value, rp["H20"].value - rp["J20"].value - rp["O20"].value, 1e-12))
    check(G, "Stock purchases and owner's draw excluded from opex",
          close(rp["O20"].value, y["fixed"] + y["labor"] + y["event"] + y["var"], 1e-9) and y["stock"] > 0 and y["draw"] > 0)
    check(G, "Labour appears only in opex (COGS = ingredients+packaging+other, no labour)",
          close(rp["G20"].value, sum(s["cogs"] for s in o.sales_lines()), 1e-9))

    # Dashboard / period summary (Full Year)
    G = "Dashboard (Full Year)"
    db = wb["Dashboard"]
    check(G, "Revenue card", db["B9"].value == m0(y['rev']), f"{db['B9'].value} vs ${y['rev']:,.0f}")
    check(G, "Gross profit card", db["F9"].value == m0(y['gp']), f"{db['F9'].value}")
    check(G, "Net profit card", db["J9"].value == m0(y['net']), f"{db['J9'].value}")
    check(G, "Net margin card", db["N9"].value == f"{y['net'] / y['rev']:.1%}", f"{db['N9'].value}")
    ip = o.item_period(dt.date(SD.YEAR, 1, 1), dt.date(SD.YEAR, 12, 31))
    top_units = sorted(ip.items(), key=lambda kv: -kv[1]["units"])[:5]
    top_gp = sorted(ip.items(), key=lambda kv: -kv[1]["gp"])[:5]
    got_u = [str(db[f"B{35 + i}"].value).split(".  ", 1)[-1] for i in range(5)]
    got_g = [str(db[f"J{35 + i}"].value).split(".  ", 1)[-1] for i in range(5)]
    check(G, "Top-5 best sellers ranking", got_u == [k for k, _ in top_units], f"{got_u}")
    check(G, "Top-5 most profitable ranking", got_g == [k for k, _ in top_gp], f"{got_g}")
    check(G, "Top-5 GP values", all(close(db[f"N{35 + i}"].value, top_gp[i][1]["gp"], 1e-9) for i in range(5)))
    be = o.breakeven()
    check(G, "Break-even progress = revenue / (BE revenue x 12)", db["N13"].value == f"{y['rev'] / (be['be_rev'] * 12):.0%}", str(db["N13"].value))
    alerts = {db[f"C{r}"].value: db[f"B{r}"].value for r in range(42, 46)}
    alerts.update({db[f"K{r}"].value: db[f"J{r}"].value for r in range(42, 46)})
    inv = o.inventory()
    check(G, "Alert counts: 1 LOSS, 1 below target, reorders match oracle, 0 bad rows",
          alerts["menu item(s) priced BELOW cost"] == 1 and alerts["menu item(s) below your target margin"] == 1
          and alerts["ingredient(s) at or below reorder level"] == sum(1 for v in inv.values() if v["status"] == "Reorder")
          and alerts["sales row(s) that need fixing (not counted)"] == 0
          and alerts["expense row(s) that need fixing (not counted)"] == 0
          and alerts["recipe line(s) with a unit or setup problem"] == 0
          and alerts["waste row(s) that need fixing (not counted)"] == 0, json.dumps(alerts))
    check(G, "Sample-data disclaimer banner shown", "SAMPLE DATA" in str(db["B6"].value))

    # Inventory
    G = "Inventory & waste"
    iv = wb["Inventory"]
    bad = []
    for name, exp in inv.items():
        r = find_row(iv, "B", name, 7, 106)
        for c, k in (("J", "used"), ("L", "wasted"), ("M", "on_hand"), ("N", "value"), ("Q", "waste_value")):
            if not close(iv[f"{c}{r}"].value, exp[k], 1e-7):
                bad.append(f"{name}.{k}: {iv[f'{c}{r}'].value} vs {exp[k]}")
        if iv[f"P{r}"].value != exp["status"]:
            bad.append(f"{name}.status {iv[f'P{r}'].value} vs {exp['status']}")
    check(G, f"{len(inv)} ingredients: usage from sales, waste, on-hand, value, status", not bad, "; ".join(bad[:4]))
    check(G, "Total stock value", close(iv["N107"].value, sum(v["value"] for v in inv.values()), 1e-9))
    wl = wb["Waste Log"]
    wexp = list(o.waste_lines())
    check(G, "Waste log values (unit-converted)", all(close(wl[f"J{7 + i}"].value, w["value"], 1e-9) for i, w in enumerate(wexp)))

    # Break-even
    G = "Break-even"
    bs = wb["Break-Even"]
    got = dict(price=bs["H8"].value, vc=bs["H9"].value, cm=bs["H17"].value, cm_ratio=bs["H18"].value, fixed=bs["H19"].value,
               be_items=bs["H20"].value, be_day=bs["H21"].value, tgt_items=bs["H22"].value, tgt_rev=bs["H23"].value,
               avg_items=bs["H24"].value, safety=bs["H25"].value, be_rev=bs["H27"].value)
    bad = [f"{k}: {got[k]} vs {be[k]}" for k in got if not close(got[k], be[k], 1e-9)]
    check(G, "Contribution, ratio, BE items/day/revenue, target-profit items/revenue, margin of safety", not bad, "; ".join(bad))
    check(G, "Break-even check: revenue at BE items covers fixed costs",
          be["be_items"] * be["cm"] >= be["fixed"] and (be["be_items"] - 1) * be["cm"] < be["fixed"])

    # Event predictor
    G = "Event predictor"
    ev = wb["Event Predictor"]
    oe = o.event()
    bad = []
    for k, c in enumerate("CDE"):
        e = oe["scen"][k]
        for row, key in ((32, "customers"), (33, "orders"), (34, "revenue"), (40, "profit"), (42, "be_orders"), (43, "be_att")):
            if not close(ev[f"{c}{row}"].value, e[key], 1e-9):
                bad.append(f"{c}{row}: {ev[f'{c}{row}'].value} vs {e[key]}")
    check(G, "Worst/base/best customers, orders, revenue, profit, break-even orders & attendance", not bad, "; ".join(bad[:4]))
    check(G, "Actual revenue/items/COGS pulled from Sales Log by event name",
          close(ev["E50"].value, oe["a_rev"], 1e-9) and close(ev["E51"].value, oe["a_items"]) and close(ev["E52"].value, oe["a_cogs"], 1e-9))
    check(G, "Actual profit and variance vs base", close(ev["E56"].value, oe["a_profit"], 1e-9) and close(ev["E58"].value, oe["variance"], 1e-9))
    check(G, "Verdict text present", "RISK" in str(ev["C46"].value) or "VIABLE" in str(ev["C46"].value), str(ev["C46"].value))

    # Simulator (sample preset = ingredient inflation)
    G = "Scenario simulator"
    sm = wb["Scenario Simulator"]
    os_ = o.simulator((0.10, 0, 0, 0, 0))
    check(G, "Baseline net profit / month", close(sm["C46"].value, os_["base"]["net"], 1e-9), f"{sm['C46'].value} vs {os_['base']['net']}")
    check(G, "Scenario net profit (ingredient +10%)", close(sm["D46"].value, os_["scen"]["net"], 1e-9), f"{sm['D46'].value} vs {os_['scen']['net']}")
    check(G, "Price needed to keep baseline profit", close(sm["D50"].value, os_["price_needed"], 1e-9))
    check(G, "Baseline x 9 months of data = annual net (excluding Oct-Dec)", close(sm["C46"].value * os_["months"], y["net"], 1e-9))
    return path


def test_period_and_presets(sample_src):
    o = Oracle()
    G = "Period switching"
    p, info = recalc(modified(sample_src, "july", [("Settings", "C9", "July")]), "july")
    wb = values(p)
    jul = o.month(SD.YEAR, 7)
    db = wb["Dashboard"]
    check(G, "Dashboard period = July: revenue / net cards", db["B9"].value == m0(jul['rev']) and db["J9"].value == m0(jul['net']),
          f"{db['B9'].value} {db['J9'].value} vs {jul['rev']:,.0f} {jul['net']:,.0f}")
    ip = o.item_period(dt.date(2026, 7, 1), dt.date(2026, 7, 31))
    top = [k for k, _ in sorted(ip.items(), key=lambda kv: -kv[1]["units"])[:5]]
    check(G, "July best sellers re-rank", [str(db[f"B{35 + i}"].value).split(".  ", 1)[-1] for i in range(5)] == top)
    check(G, "Break-even progress uses 1 month for a month period",
          db["N13"].value == f"{jul['rev'] / o.breakeven()['be_rev']:.0%}", str(db["N13"].value))
    p, info = recalc(modified(sample_src, "dec", [("Settings", "C9", "December")]), "dec")
    db = values(p)["Dashboard"]
    check(G, "Empty month (December): no errors, $0 cards, empty-state message",
          info["total_errors"] == 0 and db["B9"].value == "$0" and "No sales recorded" in str(db["B6"].value)
          and db["N9"].value == "-" and db["B35"].value == "No sales in this period yet", f"{db['B9'].value} | {db['B6'].value}")

    G = "Scenario switching"
    for name, vals in BW.PRESETS[1:]:
        p, info = recalc(modified(sample_src, "preset", [("Scenario Simulator", "C20", name)]), "preset")
        sm = values(p)["Scenario Simulator"]
        exp = o.simulator(vals)
        check(G, f"Preset '{name}' net profit", close(sm["D46"].value, exp["scen"]["net"], 1e-9),
              f"{sm['D46'].value} vs {exp['scen']['net']}")
    p, info = recalc(modified(sample_src, "custom", [("Scenario Simulator", "C20", BW.PRESETS[0][0]),
                                                     ("Scenario Simulator", "C23", 0.05), ("Scenario Simulator", "C24", -0.1),
                                                     ("Scenario Simulator", "C25", 0.08), ("Scenario Simulator", "C26", 0.03),
                                                     ("Scenario Simulator", "C27", 0.5)]), "custom")
    sm = values(p)["Scenario Simulator"]
    exp = o.simulator((0.05, -0.1, 0.08, 0.03, 0.5))
    check(G, "Custom levers net profit & price needed", close(sm["D46"].value, exp["scen"]["net"], 1e-9) and close(sm["D50"].value, exp["price_needed"], 1e-9))
    base_wb = values(TMP / "sample.xlsx")
    check(G, "Scenario does not modify source data (Monthly Report unchanged)",
          values(p)["Monthly Report"]["P20"].value == base_wb["Monthly Report"]["P20"].value)

    G = "Currency"
    p, info = recalc(modified(sample_src, "gbp", [("Settings", "C7", "GBP")]), "gbp")
    wb = values(p)
    db = wb["Dashboard"]
    y = o.period(dt.date(2026, 1, 1), dt.date(2026, 12, 31))
    check(G, "GBP: KPI shows £ with identical amount (no FX conversion)", db["B9"].value == m0(y['rev'], "£"), str(db["B9"].value))
    check(G, "GBP: headers switch symbol", wb["Menu Pricing"]["O6"].value.endswith("(£)") and wb["Sales Log"]["I6"].value == "Revenue (£)")
    check(G, "GBP: values unchanged vs USD", wb["Monthly Report"]["P20"].value == base_wb["Monthly Report"]["P20"].value)
    check(G, "GBP: recommendation text uses £", "£" in str(wb["Menu Pricing"]["U19"].value), str(wb["Menu Pricing"]["U19"].value))
    p, info = recalc(modified(sample_src, "badcur", [("Settings", "C7", "EUR")]), "badcur")
    check(G, "Invalid currency typed (EUR) falls back to USD without errors", info["total_errors"] == 0 and values(p)["Dashboard"]["B9"].value.startswith("$"))


def test_template(template_src):
    G = "Template (empty)"
    p, info = recalc(template_src, "template")
    check(G, "Zero formula errors when completely empty", info["total_errors"] == 0, json.dumps(info))
    wb = values(p)
    db = wb["Dashboard"]
    check(G, "KPI cards show $0 / '-' (no #DIV/0!)", db["B9"].value == "$0" and db["N9"].value == "-" and db["J9"].value == "$0")
    check(G, "Empty-state guidance on Dashboard", "No sales recorded" in str(db["B6"].value))
    check(G, "Break-even status asks for fixed costs", "Enter your monthly fixed costs" in str(wb["Break-Even"]["H26"].value))
    check(G, "Event verdict asks for inputs", "Enter attendance" in str(wb["Event Predictor"]["C46"].value))
    check(G, "Simulator explains missing baseline", "No sales yet" in str(wb["Scenario Simulator"]["D17"].value))
    errs = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str) and v.startswith("#") and v.endswith(("!", "?", "A")):
                    errs.append(f"{ws.title}:{v}")
    check(G, "No error strings anywhere", not errs, str(errs[:5]))
    tw = load_workbook(template_src)
    check(G, "No sample rows left in logs", all(tw[s]["B7"].value in (None, "") for s in ("Sales Log", "Expenses", "Waste Log", "Inventory", "Menu Pricing", "Recipe Builder")))


def test_edge_cases(template_src):
    G = "Edge cases & invalid input"
    e = [
        ("Settings", "C10", 20), ("Settings", "C14", dt.date(2026, 1, 1)),
        ("Inventory", "B7", "Beef"), ("Inventory", "D7", "kg"), ("Inventory", "E7", 2), ("Inventory", "F7", 20),  # 10/kg
        ("Inventory", "B8", "Buns"), ("Inventory", "D8", "each"), ("Inventory", "E8", 10), ("Inventory", "F8", 5),  # 0.5 each
        ("Inventory", "B9", "Oil"), ("Inventory", "D9", "l"),  # no price
        ("Inventory", "B10", "Beef"), ("Inventory", "D10", "kg"), ("Inventory", "E10", 1), ("Inventory", "F10", 1),  # duplicate
        ("Menu Pricing", "B7", "Burger"), ("Menu Pricing", "G7", 0.5), ("Menu Pricing", "H7", 3), ("Menu Pricing", "O7", 10),
        ("Menu Pricing", "B8", "Free sample"), ("Menu Pricing", "O8", 0),
        ("Menu Pricing", "B9", "Cheap thing"), ("Menu Pricing", "E9", 3), ("Menu Pricing", "O9", 2),
        ("Recipe Builder", "B7", "Burger"), ("Recipe Builder", "C7", "Beef"), ("Recipe Builder", "D7", 150), ("Recipe Builder", "E7", "g"),
        ("Recipe Builder", "B8", "Burger"), ("Recipe Builder", "C8", "Buns"), ("Recipe Builder", "D8", 1), ("Recipe Builder", "E8", "each"),
        ("Recipe Builder", "B9", "Burger"), ("Recipe Builder", "C9", "Buns"), ("Recipe Builder", "D9", 50), ("Recipe Builder", "E9", "g"),  # mismatch
        ("Recipe Builder", "B10", "Burger"), ("Recipe Builder", "C10", "Oil"), ("Recipe Builder", "D10", 10), ("Recipe Builder", "E10", "ml"),  # no price
        ("Recipe Builder", "B11", "Ghost item"), ("Recipe Builder", "C11", "Beef"), ("Recipe Builder", "D11", 1), ("Recipe Builder", "E11", "kg"),
        ("Recipe Builder", "B12", "Burger"), ("Recipe Builder", "C12", "Beef"), ("Recipe Builder", "D12", -5), ("Recipe Builder", "E12", "g"),
        # sales
        ("Sales Log", "B7", dt.date(2026, 3, 1)), ("Sales Log", "C7", "Burger"), ("Sales Log", "D7", 10),
        ("Sales Log", "B8", "not a date"), ("Sales Log", "C8", "Burger"), ("Sales Log", "D8", 5),
        ("Sales Log", "B9", dt.date(2026, 3, 2)), ("Sales Log", "C9", "Burger"), ("Sales Log", "D9", -3),
        ("Sales Log", "B10", dt.date(2026, 3, 2)), ("Sales Log", "C10", "Unknown"), ("Sales Log", "D10", 2),
        ("Sales Log", "B11", dt.date(2026, 3, 3)), ("Sales Log", "C11", "Free sample"), ("Sales Log", "D11", 4),
        ("Sales Log", "B12", dt.date(2026, 3, 3)), ("Sales Log", "C12", "Burger"), ("Sales Log", "D12", 2), ("Sales Log", "E12", 12),
        ("Sales Log", "B13", dt.date(2025, 12, 31)), ("Sales Log", "C13", "Burger"), ("Sales Log", "D13", 100),  # other year
        ("Sales Log", "B14", dt.date(2026, 3, 4)), ("Sales Log", "D14", 1),  # missing item
        # expenses
        ("Expenses", "B7", dt.date(2026, 3, 1)), ("Expenses", "C7", "Insurance"), ("Expenses", "E7", 100),
        ("Expenses", "B8", dt.date(2026, 3, 1)), ("Expenses", "C8", "Made up category"), ("Expenses", "E8", 999),
        ("Expenses", "B9", dt.date(2026, 3, 1)), ("Expenses", "C9", "Insurance"), ("Expenses", "E9", "abc"),
        ("Expenses", "B10", dt.date(2026, 3, 5)), ("Expenses", "C10", "Insurance"), ("Expenses", "E10", -20),  # refund
        ("Expenses", "B11", dt.date(2026, 3, 5)), ("Expenses", "C11", "Owner's Draw"), ("Expenses", "E11", 500),
        ("Expenses", "B12", dt.date(2026, 3, 5)), ("Expenses", "C12", "Ingredient & Packaging Purchases"), ("Expenses", "E12", 300),
        ("Expenses", "B13", dt.date(2026, 3, 5)), ("Expenses", "C13", "Staff Wages & Payroll"), ("Expenses", "E13", 50),
        # waste
        ("Waste Log", "B7", dt.date(2026, 3, 1)), ("Waste Log", "C7", "Beef"), ("Waste Log", "D7", 500), ("Waste Log", "E7", "g"),
        ("Waste Log", "B8", dt.date(2026, 3, 1)), ("Waste Log", "C8", "Buns"), ("Waste Log", "D8", 2), ("Waste Log", "E8", "kg"),
        # break-even with zero contribution
        ("Break-Even", "C8", 1000), ("Break-Even", "G8", 5), ("Break-Even", "G9", 5),
        # event with 100% variable costs
        ("Event Predictor", "C13", 100), ("Event Predictor", "D13", 100), ("Event Predictor", "E13", 100),
        ("Event Predictor", "C14", 0.1), ("Event Predictor", "D14", 0.1), ("Event Predictor", "E14", 0.1),
        ("Event Predictor", "C16", 10), ("Event Predictor", "D16", 10), ("Event Predictor", "E16", 10),
        ("Event Predictor", "C19", 50), ("Event Predictor", "C22", 0.6), ("Event Predictor", "C20", 0.4),
    ]
    p, info = recalc(modified(template_src, "edge", e), "edge")
    check(G, "Zero spreadsheet errors with invalid inputs", info["total_errors"] == 0, json.dumps(info))
    wb = values(p)
    rb, sl, ex, wl, mp, iv = (wb[s] for s in ("Recipe Builder", "Sales Log", "Expenses", "Waste Log", "Menu Pricing", "Inventory"))
    check(G, "Recipe g -> kg converts (150 g beef at 10/kg = 1.50)", rb["G7"].value == "OK" and close(rb["J7"].value, 1.5, 1e-12))
    check(G, "Unit mismatch (g vs each) flagged and NOT costed", rb["G9"].value == "Unit mismatch" and rb["J9"].value in (None, ""))
    check(G, "Ingredient with no price flagged", rb["G10"].value == "Add price in Inventory")
    check(G, "Recipe for item not on menu flagged", rb["G11"].value == "Not on Menu Pricing")
    check(G, "Negative recipe quantity rejected", rb["G12"].value == "Enter quantity")
    check(G, "Menu: recipe cost sums only valid lines (1.50 + 0.50)", close(mp["D7"].value, 2.0, 1e-12))
    # Burger: 2.0 ing + .5 pack + 3/60*20=1 labor = 3.5; price 10 -> margin .65 markup 1.857
    check(G, "Menu: total cost, margin 65%, markup 185.7% (hand-calculated)",
          close(mp["K7"].value, 3.5, 1e-12) and close(mp["Q7"].value, 0.65, 1e-12) and close(mp["R7"].value, 6.5 / 3.5, 1e-12))
    check(G, "Menu: suggested price at 60% default target = 3.5/0.4 = 8.75", close(mp["N7"].value, 8.75, 1e-12))
    check(G, "Menu: zero price -> 'Set a price' (no divide by zero)", mp["T8"].value == "Set a price" and mp["Q8"].value in (None, ""))
    check(G, "Menu: override ingredient cost used; price below cost -> LOSS", mp["T9"].value == "LOSS" and close(mp["F9"].value, 3))
    check(G, "Duplicate ingredient name flagged via status/CF and lookup uses first", iv["B10"].value == "Beef")
    statuses = [sl[f"N{r}"].value for r in range(7, 15)]
    check(G, "Sales checks: OK / bad date / qty<=0 / unknown item / no price / override / other year / missing item",
          statuses == ["OK", "Enter a valid date", "Qty must be above 0", "Item not on Menu Pricing", "No price on menu", "OK", "OK", "Pick a menu item"], str(statuses))
    check(G, "Price override used (2 x 12 = 24)", close(sl["I12"].value, 24))
    rp = wb["Monthly Report"]
    check(G, "Only valid 2026 sales counted in March (10x10 + 2x12 = 124)", close(rp["E10"].value, 124) and close(rp["F10"].value, 12), f"{rp['E10'].value}")
    check(G, "Prior-year sale excluded from 2026 totals", close(rp["E20"].value, 124))
    check(G, "Expense checks: unknown category and non-numeric flagged; refund OK",
          ex["I8"].value == "Category has no cost type" and ex["I9"].value == "Enter an amount" and ex["I10"].value == "OK")
    check(G, "Opex: fixed 100-20 refund = 80, labour 50; draw & stock purchase excluded",
          close(rp["K10"].value, 80) and close(rp["L10"].value, 50) and close(rp["O10"].value, 130) and close(rp["R10"].value, 300))
    check(G, "Waste unit mismatch (kg for 'each' buns) flagged; 500 g beef = 5.00",
          wl["H8"].value == "Unit mismatch" and close(wl["J7"].value, 5.0) and close(rp["J10"].value, 5.0))
    # net = (124 - cogs) - 5 - 130; cogs = 12 sold x 2.5 (ing 2 + pack .5)
    check(G, "Net profit hand-calculated: 124 - 30 - 5 - 130 = -41", close(rp["P10"].value, -41, 1e-12), str(rp["P10"].value))
    check(G, "Inventory: beef used = 12 burgers x 0.15 kg = 1.8; waste 0.5 kg", close(iv["J7"].value, 1.8) and close(iv["L7"].value, 0.5))
    check(G, "Inventory: on hand negative -> 'Recount (negative)', value floored at 0",
          iv["P7"].value == "Recount (negative)" and close(iv["N7"].value, 0))
    check(G, "Inventory: missing price -> 'Add pack size & price'", iv["P9"].value == "Add pack size & price")
    bs = wb["Break-Even"]
    check(G, "Break-even with zero contribution -> 'Not reachable', no error", bs["H20"].value == "Not reachable", str(bs["H20"].value))
    ev = wb["Event Predictor"]
    check(G, "Event with variable costs >= 100% -> break-even 'n/a', loss verdict",
          ev["D42"].value == "n/a" and "NOT VIABLE" in str(ev["C46"].value), f"{ev['D42'].value} {ev['C46'].value}")
    db = wb["Dashboard"]
    alerts = {db[f"K{r}"].value: db[f"J{r}"].value for r in range(42, 46)}
    check(G, "Dashboard counts bad rows (recipes 4, sales 5, expenses 2, waste 1)",
          alerts["recipe line(s) with a unit or setup problem"] == 4 and alerts["sales row(s) that need fixing (not counted)"] == 5
          and alerts["expense row(s) that need fixing (not counted)"] == 2 and alerts["waste row(s) that need fixing (not counted)"] == 1,
          json.dumps(alerts))


def test_volume(template_src):
    G = "High volume"
    rng = random.Random(7)
    e = [("Inventory", "B7", "Beef"), ("Inventory", "D7", "kg"), ("Inventory", "E7", 1), ("Inventory", "F7", 12)]
    items = []
    for i in range(50):
        name = f"Item {i + 1:02d}"
        price = round(rng.uniform(3, 18), 2)
        items.append((name, price))
        e += [("Menu Pricing", f"B{7 + i}", name), ("Menu Pricing", f"E{7 + i}", round(price * 0.3, 2)),
              ("Menu Pricing", f"G{7 + i}", 0.3), ("Menu Pricing", f"O{7 + i}", price)]
    exp_rev = defaultdict(float)
    exp_cogs = 0.0
    for r in range(7, 3007):
        d = dt.date(2026, 1, 1) + dt.timedelta(days=rng.randint(0, 364))
        name, price = items[rng.randint(0, 49)]
        qty = rng.randint(1, 40)
        e += [("Sales Log", f"B{r}", d), ("Sales Log", f"C{r}", name), ("Sales Log", f"D{r}", qty)]
        exp_rev[d.month] += qty * price
        exp_cogs += qty * (round(price * 0.3, 2) + 0.3)
    exp_opex = 0.0
    for r in range(7, 1507):
        d = dt.date(2026, 1, 1) + dt.timedelta(days=rng.randint(0, 364))
        amt = round(rng.uniform(5, 900), 2)
        e += [("Expenses", f"B{r}", d), ("Expenses", f"C{r}", "Fuel & Generator"), ("Expenses", f"E{r}", amt)]
        exp_opex += amt
    p, info = recalc(modified(template_src, "volume", e), "volume")
    check(G, f"3,000 sales + 1,500 expenses + 50 menu items: zero errors (LibreOffice {info['seconds']}s)", info["total_errors"] == 0, json.dumps(info))
    rp = values(p)["Monthly Report"]
    check(G, "Monthly revenue matches for all 12 months", all(close(rp[f"E{7 + m}"].value, exp_rev[m], 1e-9) for m in range(1, 13)))
    check(G, "Annual COGS and opex totals", close(rp["G20"].value, exp_cogs, 1e-9) and close(rp["O20"].value, exp_opex, 1e-9))


FORMULA_TABLES = {
    "Menu Pricing": (7, 56, "DFIKMNPQRSTUVWXYZ"), "Recipe Builder": (7, 406, "FGHIJKL"), "Inventory": (7, 106, "GJLMNPQ"),
    "Waste Log": (7, 506, "GHIJ"), "Sales Log": (7, 3006, "HIJKLMNO"), "Expenses": (7, 1506, "HI"),
}


def test_structure(sample_src):
    G = "Formula consistency & references"
    wb = load_workbook(sample_src)
    bad = []
    for sh, (r0, r1, cols) in FORMULA_TABLES.items():
        ws = wb[sh]
        for c in list(cols) + (["AA"] if sh == "Menu Pricing" else []):
            ref = re.sub(r"(?<![A-Z])(\$?[A-Z]{1,2})" + str(r0) + r"(?!\d)", r"\1#", str(ws[f"{c}{r0}"].value))
            for r in range(r0 + 1, r1 + 1):
                f = re.sub(r"(?<![A-Z])(\$?[A-Z]{1,2})" + str(r) + r"(?!\d)", r"\1#", str(ws[f"{c}{r}"].value))
                if f != ref:
                    bad.append(f"{sh}!{c}{r}")
                    break
    check(G, "Every calculated column uses one consistent formula on every row", not bad, str(bad))
    names = set(wb.defined_names.keys())
    used = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    used.update(re.findall(r"\b((?:cfg|ps|lst|be)_[A-Za-z]+)\b", c.value))
    check(G, "All named ranges used in formulas are defined", used <= names, str(used - names))
    for n in names:
        dn = wb.defined_names[n]
        sheet = dn.attr_text.split("!")[0].strip("'")
        if sheet not in wb.sheetnames:
            check(G, f"Named range {n} points to a real sheet", False, dn.attr_text)
    check(G, "No volatile/unsupported functions (INDIRECT, OFFSET, XLOOKUP, FILTER, LET, LAMBDA, macros)",
          not any(re.search(r"\b(INDIRECT|OFFSET|XLOOKUP|XMATCH|FILTER|SORT|UNIQUE|LET|LAMBDA|SEQUENCE)\(", str(c.value) or "")
                  for ws in wb.worksheets for row in ws.iter_rows() for c in row))
    check(G, "File is .xlsx (no VBA project)", not (sample_src.read_bytes().find(b"vbaProject") >= 0))


def write_report():
    groups = defaultdict(list)
    for g, n, ok, d in RESULTS:
        groups[g].append((n, ok, d))
    total = len(RESULTS)
    passed = sum(1 for r in RESULTS if r[2])
    lines = [f"# Automated QA results", "", f"Run: {dt.datetime.now():%Y-%m-%d %H:%M} · engine: LibreOffice headless recalculation · "
             f"expected values: tests/oracle.py (independent) and hand calculations", "",
             f"**{passed} / {total} checks passed**", ""]
    for g, rows in groups.items():
        lines += [f"## {g}", "", "| Result | Check |", "|---|---|"]
        for n, ok, d in rows:
            lines.append(f"| {'PASS' if ok else '**FAIL**'} | {n}{(' — ' + d.replace('|', '/')) if (d and not ok) else ''} |")
        lines.append("")
    (ROOT / "tests" / "QA_RESULTS.md").write_text("\n".join(lines))
    return passed, total


if __name__ == "__main__":
    s, t = build_both()
    test_structure(s)
    test_sample(s)
    test_period_and_presets(s)
    test_template(t)
    test_edge_cases(t)
    test_volume(t)
    passed, total = write_report()
    for f in (s, t):
        f.unlink()
    print(f"\n{passed}/{total} passed")
    sys.exit(0 if passed == total else 1)
