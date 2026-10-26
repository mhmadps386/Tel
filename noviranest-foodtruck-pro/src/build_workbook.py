"""Build the NoviraNest Food Truck Profit & Pricing System PRO workbook.

Usage:
    python build_workbook.py sample   out.xlsx
    python build_workbook.py template out.xlsx
"""
import datetime as dt
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

import sample_data as SD
from theme import *  # noqa: F401,F403

VERSION = "1.0.0"
RELEASE_DATE = "2026-10-26"

# Sheet names --------------------------------------------------------------
START, DASH, SET, MENU, REC, INV, WASTE, SALES, EXP, REP, BE, EVT, SIM, LST = (
    "Start Here", "Dashboard", "Settings", "Menu Pricing", "Recipe Builder", "Inventory", "Waste Log",
    "Sales Log", "Expenses", "Monthly Report", "Break-Even", "Event Predictor", "Scenario Simulator", "Lists")
NAV = [START, DASH, SET]

# Table capacities -----------------------------------------------------------
FIRST = 7
N_MENU, N_ING, N_REC, N_WASTE, N_SALES, N_EXP = 50, 100, 400, 500, 3000, 1500
MENU_LAST = FIRST + N_MENU - 1
ING_LAST = FIRST + N_ING - 1
REC_LAST = FIRST + N_REC - 1
WASTE_LAST = FIRST + N_WASTE - 1
SALES_LAST = FIRST + N_SALES - 1
EXP_LAST = FIRST + N_EXP - 1


def q(sheet):
    return f"'{sheet}'"


def rng(sheet, c, first, last):
    return f"{q(sheet)}!${c}${first}:${c}${last}"


# Frequently used ranges
S_DATE, S_ITEM, S_QTY = rng(SALES, "B", FIRST, SALES_LAST), rng(SALES, "C", FIRST, SALES_LAST), rng(SALES, "D", FIRST, SALES_LAST)
S_LOC = rng(SALES, "G", FIRST, SALES_LAST)
S_REV, S_COGS, S_GP = rng(SALES, "I", FIRST, SALES_LAST), rng(SALES, "K", FIRST, SALES_LAST), rng(SALES, "L", FIRST, SALES_LAST)
S_ING, S_STATUS = rng(SALES, "M", FIRST, SALES_LAST), rng(SALES, "N", FIRST, SALES_LAST)
E_DATE, E_CAT, E_AMT, E_TYPE = (rng(EXP, c, FIRST, EXP_LAST) for c in "BCEH")
W_DATE, W_ING, W_SU, W_VAL = (rng(WASTE, c, FIRST, WASTE_LAST) for c in "BCIJ")
M_NAME = rng(MENU, "B", FIRST, MENU_LAST)
I_NAME, I_UNIT, I_CPU = (rng(INV, c, FIRST, ING_LAST) for c in "BDG")
U_NAME, U_DIM, U_FAC = (rng(LST, c, 6, 30) for c in "BCD")
CAT_NAME, CAT_TYPE = rng(LST, "F", 6, 35), rng(LST, "G", 6, 35)

MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September",
          "October", "November", "December"]


# ---------------------------------------------------------------------------
# Generic helpers
# ---------------------------------------------------------------------------
def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def dv_list(ws, source, cells, prompt=None, title=None, strict=True):
    dv = DataValidation(type="list", formula1=source, allow_blank=True)
    dv.error = "Please pick a value from the dropdown list."
    dv.errorTitle = "Not in list"
    dv.showErrorMessage = strict
    dv.errorStyle = "stop" if strict else "warning"
    if prompt:
        dv.prompt, dv.promptTitle, dv.showInputMessage = prompt, title or "Tip", True
    ws.add_data_validation(dv)
    dv.add(cells)
    return dv


def dv_number(ws, cells, minimum=0, maximum=None, prompt=None, title=None, kind="decimal", greater=False):
    if maximum is None:
        op, f1, f2 = ("greaterThan" if greater else "greaterThanOrEqual"), str(minimum), None
    else:
        op, f1, f2 = "between", str(minimum), str(maximum)
    dv = DataValidation(type=kind, operator=op, formula1=f1, formula2=f2, allow_blank=True)
    dv.error = (f"Enter a number {'greater than' if greater else 'of at least'} {minimum}"
                + (f" and at most {maximum}" if maximum is not None else "") + ".")
    dv.errorTitle = "Invalid number"
    dv.showErrorMessage = True
    if prompt:
        dv.prompt, dv.promptTitle, dv.showInputMessage = prompt, title or "Tip", True
    ws.add_data_validation(dv)
    dv.add(cells)


def dv_date(ws, cells, prompt="Enter a date, e.g. 15/03/2026 or 3/15/2026 (your computer's format)."):
    dv = DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)", formula2="DATE(2100,12,31)",
                        allow_blank=True)
    dv.error, dv.errorTitle, dv.showErrorMessage = "Enter a valid date (2000-2100).", "Invalid date", True
    dv.prompt, dv.promptTitle, dv.showInputMessage = prompt, "Date", True
    ws.add_data_validation(dv)
    dv.add(cells)


def headers(ws, row, cols_titles, height=34):
    ws.row_dimensions[row].height = height
    for c, t in cols_titles.items():
        style_header(ws[f"{c}{row}"])
        ws[f"{c}{row}"] = t


def group_band(ws, row, start, end, text, color=TEAL):
    ws.merge_cells(f"{start}{row}:{end}{row}")
    style_group(ws[f"{start}{row}"], color)
    ws[f"{start}{row}"] = text
    ws.row_dimensions[row].height = 16


STATUS_COLS = {"Menu Pricing": "T", "Recipe Builder": "G", "Inventory": "P", "Waste Log": "H", "Sales Log": "N", "Expenses": "I"}


def table_rows(ws, first, last, inputs, calcs, formulas, formats, height=17):
    """inputs/calcs: column letters. formulas: col -> template using {r}. formats: col -> fmt."""
    sc = STATUS_COLS.get(ws.title)
    for r in range(first, last + 1):
        ws.row_dimensions[r].height = height
        for c in inputs:
            style_input(ws[f"{c}{r}"], formats.get(c))
        for c in calcs:
            cell = ws[f"{c}{r}"]
            style_calc(cell, formats.get(c))
            if c in formulas:
                cell.value = formulas[c].format(r=r)
        if sc:
            ws[f"{sc}{r}"].alignment = Alignment(horizontal="center", vertical="center")


def info_box(ws, cell_range, text, height=None):
    ws.merge_cells(cell_range)
    first = cell_range.split(":")[0]
    style_info(ws[first])
    ws[first] = text
    ws[first].alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    if height:
        r = int("".join(ch for ch in first if ch.isdigit()))
        ws.row_dimensions[r].height = height


def legend(ws, row, col_letter="B"):
    """Small colour legend: Input / Calculated / Warning."""
    from openpyxl.utils import column_index_from_string as ci
    c0 = ci(col_letter)
    items = [("Type here", "input"), ("Automatic", "calc"), ("Needs attention", "warn")]
    for i, (txt, kind) in enumerate(items):
        cell = ws.cell(row=row, column=c0 + i)
        cell.value = txt
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if kind == "input":
            style_input(cell, align="center")
            cell.protection = cell.protection.copy(locked=True)
        elif kind == "calc":
            style_calc(cell, align="center")
        else:
            cell.fill, cell.font = fill(WARN_FILL), font(10, True, WARN_FONT)
            cell.border = BOX
        cell.font = font(8, bold=True, color=cell.font.color.rgb if cell.font.color else INK)


def money_fmt(ws, ranges, whole=False):
    currency_cf(ws, ranges, "cfg_Currency", whole)


def sym_header(text):
    """Header text that shows the active currency symbol, e.g. 'Menu price ($)'."""
    return f'="{text} ("&cfg_Symbol&")"'


def print_setup(ws, area, landscape=True, fit_width=True):
    ws.print_area = area
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5


# ---------------------------------------------------------------------------
# Lists
# ---------------------------------------------------------------------------
def build_lists(wb, ws):
    sheet_banner(ws, "Lists & Units", "Dropdown sources and unit conversion factors. Add your own categories in the empty cream rows.",
                 20, NAV, tab_color=MUTED)
    set_widths(ws, {"B": 14, "C": 11, "D": 13, "E": 3, "F": 32, "G": 24, "H": 40, "I": 3, "J": 18, "K": 3,
                    "L": 16, "M": 3, "N": 18, "O": 3, "P": 14, "Q": 3, "R": 10, "S": 3, "T": 14})
    headers(ws, 5, {"B": "Unit", "C": "Dimension", "D": "Factor to base\n(g / ml / each)", "F": "Expense category",
                    "G": "Cost type", "H": "What goes here", "J": "Payment method", "L": "Menu category",
                    "N": "Waste reason", "P": "Dashboard period", "R": "Currency", "T": "Price rounding step"})
    ws.row_dimensions[5].height = 30
    for i in range(25):
        r = 6 + i
        for c in "BCD":
            style_input(ws[f"{c}{r}"], QTY if c == "D" else None)
        if i < len(SD.UNITS):
            u, d, f = SD.UNITS[i]
            ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = u, d, f
    for i in range(30):
        r = 6 + i
        for c in "FGH":
            style_input(ws[f"{c}{r}"])
        if i < len(SD.EXPENSE_CATEGORIES):
            n, t, note = SD.EXPENSE_CATEGORIES[i]
            ws[f"F{r}"], ws[f"G{r}"], ws[f"H{r}"] = n, t, note
    dv_list(ws, "=$V$6:$V$11", f"G6:G35", "Choose how this category is treated in profit.", "Cost type")
    dv_list(ws, '"Mass,Volume,Count"', "C6:C30")
    for i, t in enumerate(SD.COST_TYPES):
        ws[f"V{6 + i}"] = t
        style_calc(ws[f"V{6 + i}"])
    ws["V5"] = "Cost types"
    style_header(ws["V5"])
    ws.column_dimensions["U"].width = 3
    ws.column_dimensions["V"].width = 26
    simple = {"J": SD.PAYMENT_METHODS, "L": SD.MENU_CATEGORIES, "N": SD.WASTE_REASONS}
    for c, values in simple.items():
        for i in range(15):
            style_input(ws[f"{c}{6 + i}"])
            if i < len(values):
                ws[f"{c}{6 + i}"] = values[i]
    for i, v in enumerate(["Full Year"] + MONTHS):
        ws[f"P{6 + i}"] = v
        style_calc(ws[f"P{6 + i}"])
    for i, v in enumerate(["USD", "GBP"]):
        ws[f"R{6 + i}"] = v
        style_calc(ws[f"R{6 + i}"])
    for i, v in enumerate([0, 0.05, 0.10, 0.25, 0.50, 1.00]):
        ws[f"T{6 + i}"] = v
        style_calc(ws[f"T{6 + i}"], "0.00")
    info_box(ws, "B33:D36",
             "Units convert only within the same dimension (Mass, Volume, Count). Weight-to-volume needs a density "
             "and is never guessed: the line is flagged 'Unit mismatch' instead.")
    info_box(ws, "F38:H41",
             "Cost types: Fixed, Labor, Event Fees and Variable count as operating expenses. 'Excluded' types are "
             "recorded for your cash records but NOT deducted from profit (food used is already in COGS; owner's "
             "draw is not a business expense).")
    ws.freeze_panes = "A6"
    print_setup(ws, "A1:V41")
    add_name(wb, "lst_Units", f"{q(LST)}!$B$6:$B$30")
    add_name(wb, "lst_ExpCat", f"{q(LST)}!$F$6:$F$35")
    add_name(wb, "lst_Pay", f"{q(LST)}!$J$6:$J$20")
    add_name(wb, "lst_MenuCat", f"{q(LST)}!$L$6:$L$20")
    add_name(wb, "lst_Waste", f"{q(LST)}!$N$6:$N$20")
    add_name(wb, "lst_Period", f"{q(LST)}!$P$6:$P$18")


# ---------------------------------------------------------------------------
# Settings
# ---------------------------------------------------------------------------
def build_settings(wb, ws, sample):
    sheet_banner(ws, "Settings", "Set these once. Every sheet reads from here.", 6, [START, DASH, MENU])
    set_widths(ws, {"B": 38, "C": 46, "D": 3, "E": 74, "F": 3})
    headers(ws, 5, {"B": "Setting", "C": "Your value", "E": "What it does"}, 20)
    rows = [
        (6, "Business name", SD.BUSINESS if sample else "Your Food Truck Name", None,
         "Shown on the Dashboard title."),
        (7, "Currency", "USD", None,
         "USD ($) or GBP (£). Changes symbols and formats only - amounts are NOT converted."),
        (8, "Reporting year", SD.YEAR if sample else 2026, "0",
         "The year shown in the Monthly Report and Dashboard charts."),
        (9, "Dashboard period", "Full Year", None,
         "Full Year or a single month of the reporting year. KPI cards and item rankings follow this."),
        (10, "Labour cost per hour (incl. payroll taxes)", SD.LABOR_RATE if sample else None, MONEY,
         "Used only to price prep time on Menu Pricing. Real wages go in Expenses (no double counting)."),
        (11, "Default target profit margin", SD.TARGET_MARGIN if sample else 0.60, PCT,
         "Margin = profit / price. Used for suggested prices unless an item has its own target."),
        (12, "Target food cost %", SD.FOODCOST_TARGET if sample else 0.32, PCT,
         "Ingredient cost / price. Items above this are flagged 'High food cost'. Typical: 25-35%."),
        (13, "Round suggested prices up to", SD.ROUND_STEP if sample else 0.25, "0.00",
         "Suggested prices are rounded UP to this step (0 = no rounding)."),
        (14, "Last stock count date", SD.COUNT_DATE if sample else None, DATE,
         "Date you last counted stock. Inventory usage and waste are counted from this date."),
    ]
    for r, label, val, fmt, help_ in rows:
        ws.row_dimensions[r].height = 22
        style_label(ws[f"B{r}"], bold=True)
        ws[f"B{r}"] = label
        style_input(ws[f"C{r}"], fmt)
        ws[f"C{r}"] = val
        style_label(ws[f"E{r}"], size=9, color=MUTED, wrap=True)
        ws[f"E{r}"] = help_
    dv_list(ws, f"={q(LST)}!$R$6:$R$7", "C7", "Choose USD or GBP.", "Currency")
    dv_list(ws, f"={q(LST)}!$P$6:$P$18", "C9", "Full Year or one month.", "Dashboard period")
    dv_list(ws, f"={q(LST)}!$T$6:$T$11", "C13", "Rounding step for suggested prices.", "Rounding")
    dv_number(ws, "C8", 2000, 2100, "Four-digit year, e.g. 2026", "Year", kind="whole")
    dv_number(ws, "C10", 0, 1000, "Hourly cost of prep labour.", "Labour rate")
    dv_number(ws, "C11", 0, 0.95, "Enter as a percentage, e.g. 60%.", "Target margin")
    dv_number(ws, "C12", 0, 1, "Enter as a percentage, e.g. 30%.", "Food cost target")
    dv_date(ws, "C14")

    group_band(ws, 17, "B", "E", "AUTOMATIC VALUES (used by formulas - do not edit)", INK)
    calc = [
        (18, "Currency symbol", '=IF($C$7="GBP","£","$")', None),
        (19, "Active currency", '=IF($C$7="GBP","GBP","USD")', None),
        (20, "Period month (0 = full year)", '=IFERROR(MATCH($C$9,lst_Period,0)-1,0)', "0"),
        (21, "Reporting year used", '=IF(AND(ISNUMBER($C$8),$C$8>=2000,$C$8<=2100),INT($C$8),2026)', "0"),
        (22, "Period start", '=IF($C$20=0,DATE($C$21,1,1),DATE($C$21,$C$20,1))', DATE),
        (23, "Period end", '=IF($C$20=0,DATE($C$21,12,31),EOMONTH($C$22,0))', DATE),
        (24, "Months in period", '=IF($C$20=0,12,1)', "0"),
        (25, "Period label", '=IF($C$20=0,"Full Year "&$C$21,$C$9&" "&$C$21)', None),
        (26, "Stock count date used", '=IF(ISNUMBER($C$14),$C$14,DATE(2000,1,1))', DATE),
        (27, "Labour rate used", '=N($C$10)', MONEY),
        (28, "Target margin used", '=IF(AND(ISNUMBER($C$11),$C$11<0.95),MAX(0,$C$11),0.6)', PCT),
        (29, "Food cost target used", '=IF(ISNUMBER($C$12),$C$12,0.32)', PCT),
        (30, "Rounding step used", '=IF(AND(ISNUMBER($C$13),$C$13>0),$C$13,0)', "0.00"),
    ]
    for r, label, f, fmt in calc:
        style_label(ws[f"B{r}"], size=9, color=MUTED)
        ws[f"B{r}"] = label
        style_calc(ws[f"C{r}"], fmt)
        ws[f"C{r}"] = f
    names = {"cfg_Business": "$C$6", "cfg_Symbol": "$C$18", "cfg_Currency": "$C$19", "cfg_Year": "$C$21",
             "cfg_Start": "$C$22", "cfg_End": "$C$23", "cfg_Months": "$C$24", "cfg_PeriodLabel": "$C$25",
             "cfg_CountDate": "$C$26", "cfg_LaborRate": "$C$27", "cfg_TargetMargin": "$C$28",
             "cfg_FoodCostTarget": "$C$29", "cfg_Round": "$C$30"}
    for n, ref in names.items():
        add_name(wb, n, f"{q(SET)}!{ref}")
    info_box(ws, "B32:E34",
             "Tip: switch Dashboard period to a month (e.g. 'July') to see that month's KPIs and best sellers. "
             "Charts on the Dashboard always show all 12 months of the reporting year.", 30)
    print_setup(ws, "A1:F34", landscape=False)


# ---------------------------------------------------------------------------
# Inventory
# ---------------------------------------------------------------------------
def build_inventory(wb, ws, data):
    sheet_banner(ws, "Inventory", "Ingredient prices, estimated stock on hand and reorder alerts. Stock falls automatically "
                 "as items sell (via Recipe Builder) and as waste is logged.", 17, [START, DASH, REC, WASTE])
    set_widths(ws, {"B": 26, "C": 18, "D": 10, "E": 10, "F": 11, "G": 11, "H": 11, "I": 11, "J": 11, "K": 11,
                    "L": 11, "M": 12, "N": 12, "O": 10, "P": 17, "Q": 12})
    group_band(ws, 5, "B", "G", "INGREDIENT & PRICE")
    group_band(ws, 5, "H", "M", "STOCK SINCE LAST COUNT (in stock units)", INK)
    group_band(ws, 5, "N", "Q", "VALUE & ALERTS")
    headers(ws, 6, {"B": "Ingredient", "C": "Supplier", "D": "Stock unit", "E": "Pack size\n(stock units)",
                    "F": sym_header("Pack price"), "G": sym_header("Cost per\nstock unit"), "H": "Opening stock\nat count",
                    "I": "Purchased\nsince count", "J": "Used by sales\n(auto)", "K": "Other usage\n(staff meals...)",
                    "L": "Wasted\n(auto)", "M": "Est. on hand", "N": sym_header("Stock value"),
                    "O": "Reorder\nlevel", "P": "Status", "Q": sym_header("Waste value")}, 40)
    ws["F6"] = '="Pack price ("&cfg_Symbol&")"'
    ws["G6"] = '="Cost per unit ("&cfg_Symbol&")"'
    f = {
        "G": '=IF($B{r}="","",IF(N($E{r})<=0,"",N($F{r})/$E{r}))',
        "J": f'=IF($B{{r}}="","",SUMIFS({rng(REC, "L", FIRST, REC_LAST)},{rng(REC, "C", FIRST, REC_LAST)},$B{{r}}))',
        "L": f'=IF($B{{r}}="","",SUMIFS({W_SU},{W_ING},$B{{r}},{W_DATE},">="&cfg_CountDate))',
        "M": '=IF($B{r}="","",N($H{r})+N($I{r})-N($J{r})-N($K{r})-N($L{r}))',
        "N": '=IF($B{r}="","",IF($G{r}="",0,MAX(0,$M{r})*$G{r}))',
        "P": '=IF($B{r}="","",IF($D{r}="","Set stock unit",IF($G{r}="","Add pack size & price",'
             'IF($M{r}<0,"Recount (negative)",IF(AND(ISNUMBER($O{r}),$M{r}<=$O{r}),"Reorder","OK")))))',
        "Q": '=IF(OR($B{r}="",$G{r}=""),"",N($L{r})*$G{r})',
    }
    fm = {"E": QTY, "F": MONEY, "G": '#,##0.0000', "H": QTY, "I": QTY, "J": '#,##0.00', "K": QTY, "L": '#,##0.00',
          "M": '#,##0.00', "N": MONEY, "O": QTY, "Q": MONEY}
    table_rows(ws, FIRST, ING_LAST, "BCDEFHIKO", "GJLMNPQ", f, fm)
    for i, ing in enumerate(data["inventory"] if data else []):
        r = FIRST + i
        ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = ing["name"], ing["supplier"], ing["unit"]
        ws[f"E{r}"], ws[f"F{r}"], ws[f"H{r}"] = ing["pack"], ing["price"], ing["opening"]
        ws[f"I{r}"], ws[f"O{r}"] = ing["purchased"], ing["reorder"]
    total = ING_LAST + 1
    ws[f"B{total}"] = "TOTAL STOCK VALUE"
    style_label(ws[f"B{total}"], bold=True)
    for c, fx in {"N": f"=SUM(N{FIRST}:N{ING_LAST})", "Q": f"=SUM(Q{FIRST}:Q{ING_LAST})"}.items():
        style_calc(ws[f"{c}{total}"], MONEY, bold=True)
        ws[f"{c}{total}"] = fx
    ws[f"P{total}"] = f'=COUNTIF(P{FIRST}:P{ING_LAST},"Reorder")&" to reorder"'
    style_calc(ws[f"P{total}"], bold=True)
    dv_list(ws, "=lst_Units", f"D{FIRST}:D{ING_LAST}", "The unit you COUNT this ingredient in (kg, lb, each...).", "Stock unit")
    dv_number(ws, f"E{FIRST}:E{ING_LAST}", 0, None, "How many stock units are in one pack (e.g. 10 for a 10 lb case).",
              "Pack size", greater=True)
    dv_number(ws, f"F{FIRST}:F{ING_LAST}", 0, None, "What you pay for one pack.", "Pack price")
    dv_number(ws, f"H{FIRST}:I{ING_LAST}", 0, None)
    dv_number(ws, f"K{FIRST}:K{ING_LAST}", 0, None)
    dv_number(ws, f"O{FIRST}:O{ING_LAST}", 0, None, "Alert when on-hand falls to this level.", "Reorder level")
    body = f"B{FIRST}:Q{ING_LAST}"
    text_rule(ws, f"P{FIRST}:P{ING_LAST}", f"P{FIRST}", "Reorder", CAUTION_FONT, CAUTION_FILL, mode="eq")
    text_rule(ws, f"P{FIRST}:P{ING_LAST}", f"P{FIRST}", "OK", OK_FONT, OK_FILL, mode="eq")
    expr_rule(ws, f"P{FIRST}:P{ING_LAST}", f'AND($P{FIRST}<>"",$P{FIRST}<>"OK",$P{FIRST}<>"Reorder")', WARN_FONT, WARN_FILL, True)
    expr_rule(ws, f"B{FIRST}:B{ING_LAST}", f'AND($B{FIRST}<>"",COUNTIF($B${FIRST}:$B${ING_LAST},$B{FIRST})>1)',
              WARN_FONT, WARN_FILL, True)
    expr_rule(ws, f"M{FIRST}:M{ING_LAST}", f'AND(ISNUMBER($M{FIRST}),$M{FIRST}<0)', WARN_FONT)
    money_fmt(ws, [f"F{FIRST}:F{ING_LAST}", f"N{FIRST}:N{total}", f"Q{FIRST}:Q{total}"])
    ws.freeze_panes = f"C{FIRST}"
    ws.auto_filter.ref = f"B6:Q{ING_LAST}"
    ws[f"B{total + 2}"] = ("Duplicate ingredient names turn red - each name must be unique. "
                           "'Used by sales' = recipe quantity x items sold since the stock count date.")
    style_label(ws[f"B{total + 2}"], size=9, color=MUTED)
    print_setup(ws, f"A1:Q{total}")


# ---------------------------------------------------------------------------
# Recipe Builder
# ---------------------------------------------------------------------------
def build_recipes(wb, ws, sample):
    sheet_banner(ws, "Recipe Builder", "One row per ingredient per menu item. Quantities convert automatically to each "
                 "ingredient's stock unit. Incompatible units are flagged and never costed.", 12,
                 [START, DASH, MENU, INV])
    set_widths(ws, {"B": 26, "C": 26, "D": 10, "E": 11, "F": 11, "G": 22, "H": 13, "I": 12, "J": 11, "K": 12, "L": 13})
    group_band(ws, 5, "B", "E", "WHAT GOES INTO ONE SERVING")
    group_band(ws, 5, "F", "L", "AUTOMATIC COSTING & USAGE", INK)
    headers(ws, 6, {"B": "Menu item", "C": "Ingredient", "D": "Qty per\nserving", "E": "Unit", "F": "Stock unit",
                    "G": "Unit check", "H": "Qty in\nstock units", "I": "", "J": "", "K": "Items sold\nsince count",
                    "L": "Usage since\ncount"}, 40)
    ws["I6"] = '="Cost per stock unit ("&cfg_Symbol&")"'
    ws["J6"] = '="Line cost ("&cfg_Symbol&")"'
    look = lambda col: f"INDEX({rng(INV, col, FIRST, ING_LAST)},MATCH($C{{r}},{I_NAME},0))"
    dim = lambda cell: f"IFERROR(INDEX({U_DIM},MATCH({cell},{U_NAME},0)),\"?\")"
    fac = lambda cell: f"INDEX({U_FAC},MATCH({cell},{U_NAME},0))"
    f = {
        "F": f'=IF($C{{r}}="","",IFERROR({look("D")}&"","?"))',
        "G": (f'=IF(AND($B{{r}}="",$C{{r}}=""),"",IF($B{{r}}="","Pick menu item",IF(ISNA(MATCH($B{{r}},{M_NAME},0)),'
              f'"Not on Menu Pricing",IF($C{{r}}="","Pick ingredient",IF($F{{r}}="?","Not in Inventory",'
              f'IF($F{{r}}="","Set stock unit in Inventory",IF(N($D{{r}})<=0,"Enter quantity",IF($E{{r}}="","Pick unit",'
              f'IF({dim("$E{r}")}<>{dim("$F{r}")},"Unit mismatch",IF(ISNUMBER({look("G")}),"OK","Add price in Inventory"))))))))))'),
        "H": f'=IF($G{{r}}<>"OK","",$D{{r}}*{fac("$E{r}")}/{fac("$F{r}")})',
        "I": f'=IF($G{{r}}<>"OK","",{look("G")})',
        "J": '=IF($G{r}<>"OK","",$H{r}*$I{r})',
        "K": f'=IF($G{{r}}<>"OK","",SUMIFS({S_QTY},{S_ITEM},$B{{r}},{S_DATE},">="&cfg_CountDate,{S_STATUS},"OK"))',
        "L": '=IF($G{r}<>"OK","",$H{r}*$K{r})',
    }
    fm = {"D": QTY, "H": '#,##0.0000', "I": '#,##0.0000', "J": '#,##0.000', "K": INT, "L": '#,##0.00'}
    table_rows(ws, FIRST, REC_LAST, "BCDE", "FGHIJKL", f, fm)
    if sample:
        r = FIRST
        for item, lines in SD.RECIPES.items():
            for ing, qty, unit in lines:
                ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"] = item, ing, qty, unit
                r += 1
    dv_list(ws, f"={rng(MENU, 'B', FIRST, MENU_LAST)}", f"B{FIRST}:B{REC_LAST}", "Pick a menu item (add items on Menu Pricing first).", "Menu item")
    dv_list(ws, f"={I_NAME}", f"C{FIRST}:C{REC_LAST}", "Pick an ingredient (add them on Inventory first).", "Ingredient")
    dv_list(ws, "=lst_Units", f"E{FIRST}:E{REC_LAST}", "Any unit of the same type as the stock unit (mass, volume or count).", "Unit")
    dv_number(ws, f"D{FIRST}:D{REC_LAST}", 0, None, "Quantity used in ONE serving.", "Quantity", greater=True)
    text_rule(ws, f"G{FIRST}:G{REC_LAST}", f"G{FIRST}", "OK", OK_FONT, OK_FILL, mode="eq")
    expr_rule(ws, f"G{FIRST}:G{REC_LAST}", f'AND($G{FIRST}<>"",$G{FIRST}<>"OK")', WARN_FONT, WARN_FILL, True)
    money_fmt(ws, [f"J{FIRST}:J{REC_LAST}"])
    ws.freeze_panes = f"C{FIRST}"
    ws.auto_filter.ref = f"B6:L{REC_LAST}"
    print_setup(ws, f"A1:L{REC_LAST}")


# ---------------------------------------------------------------------------
# Menu Pricing
# ---------------------------------------------------------------------------
def build_menu(wb, ws, sample):
    sheet_banner(ws, "Smart Menu Pricing", "Cost per serving, suggested price, real margin vs markup, and what to fix. "
                 "Sales columns follow the Dashboard period.", 25, [START, DASH, REC, SET])
    set_widths(ws, {"B": 24, "C": 11, "D": 11, "E": 11, "F": 11, "G": 10, "H": 9, "I": 10, "J": 10, "K": 11,
                    "L": 9, "M": 9, "N": 11, "O": 11, "P": 10, "Q": 9, "R": 9, "S": 9, "T": 15, "U": 40,
                    "V": 9, "W": 12, "X": 12, "Y": 9, "Z": 4, "AA": 4})
    group_band(ws, 5, "B", "C", "ITEM")
    group_band(ws, 5, "D", "K", "COST PER SERVING", INK)
    group_band(ws, 5, "L", "N", "PRICING TARGET")
    group_band(ws, 5, "O", "U", "YOUR PRICE & RESULT", AMBER)
    group_band(ws, 5, "V", "Y", "SALES (DASHBOARD PERIOD)", INK)
    hdr = {"B": "Menu item", "C": "Category", "H": "Prep\nminutes", "L": "Target\nmargin", "M": "Target\nused",
           "Q": "Actual\nmargin", "R": "Markup\non cost", "S": "Food\ncost %", "T": "Status", "U": "Recommendation",
           "V": "Items\nsold", "Y": "Sales\nmix"}
    headers(ws, 6, hdr, 44)
    for c, t in {"D": "Recipe\ningredients", "E": "Ingredient\noverride", "F": "Ingredient\ncost used",
                 "G": "Packaging", "I": "Prep\nlabour", "J": "Other\nvariable", "K": "TOTAL COST\nper serving",
                 "N": "Suggested\nprice", "O": "YOUR MENU\nPRICE", "P": "Profit per\nserving", "W": "Revenue",
                 "X": "Gross\nprofit"}.items():
        style_header(ws[f"{c}6"])
        ws[f"{c}6"] = f'="{t} ("&cfg_Symbol&")"'
    menu_ix = f"MATCH($B{{r}},{M_NAME},0)"
    f = {
        "D": f'=IF($B{{r}}="","",SUMIFS({rng(REC, "J", FIRST, REC_LAST)},{rng(REC, "B", FIRST, REC_LAST)},$B{{r}}))',
        "F": '=IF($B{r}="","",IF(ISNUMBER($E{r}),$E{r},N($D{r})))',
        "I": '=IF($B{r}="","",N($H{r})/60*cfg_LaborRate)',
        "K": '=IF($B{r}="","",$F{r}+N($G{r})+$I{r}+N($J{r}))',
        "M": '=IF($B{r}="","",IF(AND(ISNUMBER($L{r}),$L{r}<0.95),MAX(0,$L{r}),cfg_TargetMargin))',
        "N": '=IF($B{r}="","",IF(cfg_Round>0,CEILING($K{r}/(1-$M{r}),cfg_Round),ROUND($K{r}/(1-$M{r}),2)))',
        "P": '=IF(OR($B{r}="",N($O{r})<=0),"",$O{r}-$K{r})',
        "Q": '=IF(OR($B{r}="",N($O{r})<=0),"",$P{r}/$O{r})',
        "R": '=IF(OR($B{r}="",N($O{r})<=0,N($K{r})<=0),"",$P{r}/$K{r})',
        "S": '=IF(OR($B{r}="",N($O{r})<=0),"",$F{r}/$O{r})',
        "T": ('=IF($B{r}="","",IF(COUNTIF($B$7:$B$56,$B{r})>1,"Duplicate name",IF(N($O{r})<=0,"Set a price",'
              'IF(N($K{r})<=0,"Add costs",IF($P{r}<0,"LOSS",IF($Q{r}<$M{r}-0.00001,"Below target",'
              'IF($S{r}>cfg_FoodCostTarget,"High food cost","Healthy")))))))'),
        "U": ('=IF($T{r}="","",IF($T{r}="Duplicate name","Each menu item needs a unique name",'
              'IF($T{r}="Set a price","Suggested starting price: "&cfg_Symbol&TEXT($N{r},"0.00"),'
              'IF($T{r}="Add costs","Build the recipe or type an ingredient cost override",'
              'IF($T{r}="LOSS","Losing "&cfg_Symbol&TEXT(-$P{r},"0.00")&" per sale - raise to "&cfg_Symbol&TEXT($N{r},"0.00"),'
              'IF($T{r}="Below target","Raise to "&cfg_Symbol&TEXT($N{r},"0.00")&" for a "&TEXT($M{r},"0%")&" margin",'
              'IF($T{r}="High food cost","Margin OK; food cost "&TEXT($S{r},"0%")&" is above "&TEXT(cfg_FoodCostTarget,"0%")&" target",'
              '"On target")))))))'),
        "V": f'=IF($B{{r}}="","",SUMIFS({S_QTY},{S_ITEM},$B{{r}},{S_DATE},">="&cfg_Start,{S_DATE},"<="&cfg_End,{S_STATUS},"OK"))',
        "W": f'=IF($B{{r}}="","",SUMIFS({S_REV},{S_ITEM},$B{{r}},{S_DATE},">="&cfg_Start,{S_DATE},"<="&cfg_End,{S_STATUS},"OK"))',
        "X": f'=IF($B{{r}}="","",SUMIFS({S_GP},{S_ITEM},$B{{r}},{S_DATE},">="&cfg_Start,{S_DATE},"<="&cfg_End,{S_STATUS},"OK"))',
        "Y": f'=IF($B{{r}}="","",IF(SUM($V${FIRST}:$V${MENU_LAST})=0,0,$V{{r}}/SUM($V${FIRST}:$V${MENU_LAST})))',
        "Z": '=IF($B{r}="",-1,N($V{r})+ROW()/1000000)',
        "AA": '=IF($B{r}="",-1E+15,N($X{r})+ROW()/1000000)',
    }
    fm = {"D": MONEY, "E": MONEY, "F": MONEY, "G": MONEY, "H": "0.#", "I": MONEY, "J": MONEY, "K": MONEY,
          "L": PCT, "M": PCT, "N": MONEY, "O": MONEY, "P": MONEY, "Q": PCT, "R": PCT, "S": PCT,
          "V": INT, "W": MONEY, "X": MONEY, "Y": PCT, "Z": "0.000000", "AA": "0.000000"}
    table_rows(ws, FIRST, MENU_LAST, list("BCEGHJLO"), list("DFIKMNPQRSTUVWXYZ") + ["AA"], f, fm, height=19)
    for r in range(FIRST, MENU_LAST + 1):
        ws[f"K{r}"].font = font(10, bold=True)
        ws[f"O{r}"].font = font(10, bold=True, color=INPUT_FONT)
        ws[f"U{r}"].font = font(9, color=CALC_FONT)
    ws.column_dimensions["Z"].hidden = True
    ws.column_dimensions["AA"].hidden = True
    if sample:
        for i, (name, cat, pk, mins, other, tm, price, _b) in enumerate(SD.MENU):
            r = FIRST + i
            ws[f"B{r}"], ws[f"C{r}"], ws[f"G{r}"], ws[f"H{r}"] = name, cat, pk, mins
            ws[f"J{r}"], ws[f"L{r}"], ws[f"O{r}"] = other, tm, price
    total = MENU_LAST + 1
    ws[f"B{total}"] = "TOTAL / ALERTS"
    style_label(ws[f"B{total}"], bold=True)
    for c in "VWX":
        style_calc(ws[f"{c}{total}"], fm[c], bold=True)
        ws[f"{c}{total}"] = f"=SUM({c}{FIRST}:{c}{MENU_LAST})"
    style_calc(ws[f"T{total}"], bold=True)
    ws[f"T{total}"] = (f'=COUNTIF(T{FIRST}:T{MENU_LAST},"LOSS")+COUNTIF(T{FIRST}:T{MENU_LAST},"Below target")'
                       f'&" need attention"')
    dv_list(ws, "=lst_MenuCat", f"C{FIRST}:C{MENU_LAST}", strict=False)
    for c, p in {"E": "Optional: type the ingredient cost per serving if you don't use Recipe Builder.",
                 "G": "Box, cup, napkin, fork... per serving.", "J": "Condiments, gas share, card fee per item, etc.",
                 "O": "The price on your menu board."}.items():
        dv_number(ws, f"{c}{FIRST}:{c}{MENU_LAST}", 0, None, p, "Money per serving")
    dv_number(ws, f"H{FIRST}:H{MENU_LAST}", 0, 600, "Hands-on minutes to prepare one serving.", "Prep minutes")
    dv_number(ws, f"L{FIRST}:L{MENU_LAST}", 0, 0.95, "Optional item target margin (blank = Settings default).", "Target margin")
    st = f"T{FIRST}:T{MENU_LAST}"
    text_rule(ws, st, f"T{FIRST}", "LOSS", WHITE, "C0362C", mode="eq")
    text_rule(ws, st, f"T{FIRST}", "Below target", CAUTION_FONT, CAUTION_FILL, mode="eq")
    text_rule(ws, st, f"T{FIRST}", "High food cost", CAUTION_FONT, CAUTION_FILL, mode="eq")
    text_rule(ws, st, f"T{FIRST}", "Healthy", OK_FONT, OK_FILL, mode="eq")
    expr_rule(ws, st, f'OR($T{FIRST}="Set a price",$T{FIRST}="Add costs",$T{FIRST}="Duplicate name")', INFO_FONT, INFO_FILL, True)
    expr_rule(ws, f"B{FIRST}:B{MENU_LAST}", f'AND($B{FIRST}<>"",COUNTIF($B${FIRST}:$B${MENU_LAST},$B{FIRST})>1)', WARN_FONT, WARN_FILL, True)
    expr_rule(ws, f"B{FIRST}:B{MENU_LAST}", f'$T{FIRST}="LOSS"', WARN_FONT, WARN_FILL, True)
    money_fmt(ws, [f"D{FIRST}:G{MENU_LAST}", f"I{FIRST}:K{MENU_LAST}", f"N{FIRST}:P{MENU_LAST}", f"W{FIRST}:X{total}"])
    ws.freeze_panes = f"C{FIRST}"
    ws.auto_filter.ref = f"B6:Y{MENU_LAST}"
    info_box(ws, f"B{total + 2}:U{total + 3}",
             "MARGIN = profit / price.  MARKUP = profit / cost.  Example: cost 4.00, price 10.00 -> margin 60%, markup 150%. "
             "Suggested price = total cost / (1 - target margin), rounded UP. Prep labour here is a pricing allowance only - "
             "real wages are recorded in Expenses. Gross profit in the sales columns uses ingredients + packaging + other variable "
             "cost (no labour), so labour is never counted twice.", 46)
    print_setup(ws, f"A1:Y{total + 3}")


# ---------------------------------------------------------------------------
# Logs
# ---------------------------------------------------------------------------
def build_sales(wb, ws, data):
    sheet_banner(ws, "Sales Log", "Record what you sold: one row per item per day (or per order). Leave 'Price charged' "
                 "blank to use your menu price.", 15, [START, DASH, MENU])
    set_widths(ws, {"B": 13, "C": 26, "D": 8, "E": 11, "F": 15, "G": 24, "H": 11, "I": 12, "J": 11, "K": 12,
                    "L": 12, "M": 12, "N": 20, "O": 5})
    group_band(ws, 5, "B", "G", "YOU TYPE")
    group_band(ws, 5, "H", "N", "AUTOMATIC", INK)
    headers(ws, 6, {"B": "Date", "C": "Menu item", "D": "Qty", "F": "Payment method", "G": "Location / event",
                    "N": "Check", "O": "Menu\nrow"}, 40)
    for c, t in {"E": "Price charged\n(optional)", "H": "Unit price\nused", "I": "Revenue", "J": "Unit cost\n(COGS)",
                 "K": "COGS", "L": "Gross\nprofit", "M": "Ingredient\ncost"}.items():
        style_header(ws[f"{c}6"])
        ws[f"{c}6"] = f'="{t} ("&cfg_Symbol&")"'
    mi = lambda col: f"INDEX({rng(MENU, col, FIRST, MENU_LAST)},$O{{r}})"
    f = {
        "O": f'=IF($C{{r}}="","",IFERROR(MATCH($C{{r}},{M_NAME},0),""))',
        "H": f'=IF(OR($C{{r}}="",$D{{r}}=""),"",IF(ISNUMBER($E{{r}}),$E{{r}},IF($O{{r}}="","",IF(N({mi("O")})>0,{mi("O")},""))))',
        "N": ('=IF(AND($B{r}="",$C{r}="",$D{r}=""),"",IF(NOT(ISNUMBER($B{r})),"Enter a valid date",IF($C{r}="","Pick a menu item",'
              'IF($O{r}="","Item not on Menu Pricing",IF(N($D{r})<=0,"Qty must be above 0",IF($H{r}="","No price on menu","OK"))))))'),
        "I": '=IF($N{r}<>"OK","",$D{r}*$H{r})',
        "J": f'=IF($N{{r}}<>"OK","",N({mi("F")})+N({mi("G")})+N({mi("J")}))',
        "K": '=IF($N{r}<>"OK","",$D{r}*$J{r})',
        "L": '=IF($N{r}<>"OK","",$I{r}-$K{r})',
        "M": f'=IF($N{{r}}<>"OK","",$D{{r}}*N({mi("F")}))',
    }
    fm = {"B": DATE, "D": QTY, "E": MONEY, "H": MONEY, "I": MONEY, "J": MONEY, "K": MONEY, "L": MONEY, "M": MONEY}
    table_rows(ws, FIRST, SALES_LAST, "BCDEFG", "HIJKLMNO", f, fm)
    for i, s in enumerate(data["sales"] if data else []):
        r = FIRST + i
        ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = s["date"], s["item"], s["qty"]
        if s["price"] is not None:
            ws[f"E{r}"] = s["price"]
        ws[f"F{r}"], ws[f"G{r}"] = s["pay"], s["loc"]
    ws.column_dimensions["O"].hidden = True
    dv_date(ws, f"B{FIRST}:B{SALES_LAST}")
    dv_list(ws, f"={M_NAME}", f"C{FIRST}:C{SALES_LAST}", "Pick a menu item.", "Menu item")
    dv_number(ws, f"D{FIRST}:D{SALES_LAST}", 0, None, "Number of servings sold.", "Quantity", greater=True)
    dv_number(ws, f"E{FIRST}:E{SALES_LAST}", 0, None, "Only if you charged something other than the menu price (event pricing, discount).", "Price charged")
    dv_list(ws, "=lst_Pay", f"F{FIRST}:F{SALES_LAST}", strict=False)
    text_rule(ws, f"N{FIRST}:N{SALES_LAST}", f"N{FIRST}", "OK", OK_FONT, OK_FILL, False, mode="eq")
    expr_rule(ws, f"N{FIRST}:N{SALES_LAST}", f'AND($N{FIRST}<>"",$N{FIRST}<>"OK")', WARN_FONT, WARN_FILL, True)
    money_fmt(ws, [f"E{FIRST}:E{SALES_LAST}", f"H{FIRST}:M{SALES_LAST}"])
    ws.freeze_panes = f"C{FIRST}"
    ws.auto_filter.ref = f"B6:N{SALES_LAST}"
    print_setup(ws, f"A1:N{SALES_LAST}")


def build_expenses(wb, ws, data):
    sheet_banner(ws, "Expenses", "Every business cost with a date and category. Ingredient purchases and owner's draw are "
                 "kept for your records but excluded from profit (see Lists).", 10, [START, DASH, LST])
    set_widths(ws, {"B": 13, "C": 30, "D": 34, "E": 12, "F": 15, "G": 22, "H": 24, "I": 20})
    group_band(ws, 5, "B", "G", "YOU TYPE")
    group_band(ws, 5, "H", "I", "AUTOMATIC", INK)
    headers(ws, 6, {"B": "Date", "C": "Category", "D": "Description", "F": "Payment method", "G": "Paid to",
                    "H": "Cost type", "I": "Check"}, 34)
    ws["E6"] = sym_header("Amount")
    style_header(ws["E6"])
    f = {
        "H": f'=IF($C{{r}}="","",IFERROR(INDEX({CAT_TYPE},MATCH($C{{r}},{CAT_NAME},0))&"","Unknown category"))',
        "I": ('=IF(AND($B{r}="",$C{r}="",$E{r}=""),"",IF(NOT(ISNUMBER($B{r})),"Enter a valid date",IF($C{r}="","Pick a category",'
              'IF(OR($H{r}="Unknown category",$H{r}=""),"Category has no cost type",IF(NOT(ISNUMBER($E{r})),"Enter an amount","OK")))))'),
    }
    table_rows(ws, FIRST, EXP_LAST, "BCDEFG", "HI", f, {"B": DATE, "E": MONEY})
    for i, e in enumerate(data["expenses"] if data else []):
        r = FIRST + i
        ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"] = e["date"], e["cat"], e["desc"], e["amt"]
        ws[f"F{r}"], ws[f"G{r}"] = e["pay"], e["vendor"]
    dv_date(ws, f"B{FIRST}:B{EXP_LAST}")
    dv_list(ws, "=lst_ExpCat", f"C{FIRST}:C{EXP_LAST}", "Pick a category. Add your own on the Lists sheet.", "Category")
    dv_list(ws, "=lst_Pay", f"F{FIRST}:F{EXP_LAST}", strict=False)
    dv_number(ws, f"E{FIRST}:E{EXP_LAST}", -1000000, 10000000, "Amount paid. Use a negative only for refunds.", "Amount")
    text_rule(ws, f"I{FIRST}:I{EXP_LAST}", f"I{FIRST}", "OK", OK_FONT, OK_FILL, False, mode="eq")
    expr_rule(ws, f"I{FIRST}:I{EXP_LAST}", f'AND($I{FIRST}<>"",$I{FIRST}<>"OK")', WARN_FONT, WARN_FILL, True)
    expr_rule(ws, f"H{FIRST}:H{EXP_LAST}", f'LEFT($H{FIRST},8)="Excluded"', MUTED)
    money_fmt(ws, [f"E{FIRST}:E{EXP_LAST}"])
    ws.freeze_panes = f"C{FIRST}"
    ws.auto_filter.ref = f"B6:I{EXP_LAST}"
    print_setup(ws, f"A1:I{EXP_LAST}")


def build_waste(wb, ws, data):
    sheet_banner(ws, "Waste Log", "Record food you throw away. Any unit of the same type as the stock unit works "
                 "(e.g. stock in kg, waste in g).", 11, [START, DASH, INV])
    set_widths(ws, {"B": 13, "C": 26, "D": 9, "E": 11, "F": 18, "G": 11, "H": 22, "I": 13, "J": 12})
    group_band(ws, 5, "B", "F", "YOU TYPE")
    group_band(ws, 5, "G", "J", "AUTOMATIC", INK)
    headers(ws, 6, {"B": "Date", "C": "Ingredient", "D": "Qty", "E": "Unit", "F": "Reason", "G": "Stock unit",
                    "H": "Check", "I": "Qty in\nstock units"}, 34)
    ws["J6"] = sym_header("Waste value")
    style_header(ws["J6"])
    look = lambda col: f"INDEX({rng(INV, col, FIRST, ING_LAST)},MATCH($C{{r}},{I_NAME},0))"
    dim = lambda cell: f"IFERROR(INDEX({U_DIM},MATCH({cell},{U_NAME},0)),\"?\")"
    fac = lambda cell: f"INDEX({U_FAC},MATCH({cell},{U_NAME},0))"
    f = {
        "G": f'=IF($C{{r}}="","",IFERROR({look("D")}&"","?"))',
        "H": (f'=IF(AND($B{{r}}="",$C{{r}}="",$D{{r}}=""),"",IF(NOT(ISNUMBER($B{{r}})),"Enter a valid date",IF($C{{r}}="","Pick ingredient",'
              f'IF($G{{r}}="?","Not in Inventory",IF($G{{r}}="","Set stock unit in Inventory",IF(N($D{{r}})<=0,"Enter quantity",'
              f'IF($E{{r}}="","Pick unit",IF({dim("$E{r}")}<>{dim("$G{r}")},"Unit mismatch",'
              f'IF(ISNUMBER({look("G")}),"OK","Add price in Inventory")))))))))'),
        "I": f'=IF($H{{r}}<>"OK","",$D{{r}}*{fac("$E{r}")}/{fac("$G{r}")})',
        "J": f'=IF($H{{r}}<>"OK","",$I{{r}}*{look("G")})',
    }
    table_rows(ws, FIRST, WASTE_LAST, "BCDEF", "GHIJ", f, {"B": DATE, "D": QTY, "I": "#,##0.000", "J": MONEY})
    for i, w in enumerate(data["waste"] if data else []):
        r = FIRST + i
        ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"], ws[f"F{r}"] = w["date"], w["ing"], w["qty"], w["unit"], w["reason"]
    dv_date(ws, f"B{FIRST}:B{WASTE_LAST}")
    dv_list(ws, f"={I_NAME}", f"C{FIRST}:C{WASTE_LAST}", "Pick an ingredient from Inventory.", "Ingredient")
    dv_list(ws, "=lst_Units", f"E{FIRST}:E{WASTE_LAST}", "Unit you measured the waste in.", "Unit")
    dv_list(ws, "=lst_Waste", f"F{FIRST}:F{WASTE_LAST}", strict=False)
    dv_number(ws, f"D{FIRST}:D{WASTE_LAST}", 0, None, "Amount thrown away.", "Quantity", greater=True)
    text_rule(ws, f"H{FIRST}:H{WASTE_LAST}", f"H{FIRST}", "OK", OK_FONT, OK_FILL, False, mode="eq")
    expr_rule(ws, f"H{FIRST}:H{WASTE_LAST}", f'AND($H{FIRST}<>"",$H{FIRST}<>"OK")', WARN_FONT, WARN_FILL, True)
    money_fmt(ws, [f"J{FIRST}:J{WASTE_LAST}"])
    ws.freeze_panes = f"C{FIRST}"
    ws.auto_filter.ref = f"B6:J{WASTE_LAST}"
    print_setup(ws, f"A1:J{WASTE_LAST}")


# ---------------------------------------------------------------------------
# Monthly Report
# ---------------------------------------------------------------------------
REP_M0, REP_M1, REP_TOT = 8, 19, 20
PS = {}  # period summary cell addresses


def build_report(wb, ws):
    sheet_banner(ws, "Monthly Report", "Profit & loss by month for the reporting year, plus the selected-period summary "
                 "that feeds the Dashboard. Everything here is automatic.", 19, [START, DASH, SET])
    set_widths(ws, {"B": 14, "C": 4, "D": 4, "E": 12, "F": 10, "G": 12, "H": 12, "I": 9, "J": 10, "K": 11, "L": 11,
                    "M": 10, "N": 11, "O": 12, "P": 12, "Q": 9, "R": 12, "S": 12})
    ws.column_dimensions["C"].hidden = True
    ws.column_dimensions["D"].hidden = True
    group_band(ws, 6, "E", "I", "SALES & GROSS PROFIT")
    group_band(ws, 6, "J", "O", "WASTE & OPERATING EXPENSES", INK)
    group_band(ws, 6, "P", "Q", "BOTTOM LINE", AMBER)
    group_band(ws, 6, "R", "S", "INFO ONLY", "6B7280")
    headers(ws, 7, {"B": "Month", "C": "Start", "D": "End", "F": "Items\nsold", "I": "Gross\nmargin", "Q": "Net\nmargin"}, 44)
    money_cols = {"E": "Revenue", "G": "COGS", "H": "Gross\nprofit", "J": "Waste", "K": "Fixed\ncosts", "L": "Labour",
                  "M": "Event\nfees", "N": "Other\nvariable", "O": "Total\nopex", "P": "NET\nPROFIT",
                  "R": "Stock\npurchases", "S": "Ingredient\npart of COGS"}
    for c, t in money_cols.items():
        style_header(ws[f"{c}7"])
        ws[f"{c}7"] = f'="{t} ("&cfg_Symbol&")"'
    by_date = lambda d1, d2: f'{S_DATE},">="&{d1},{S_DATE},"<="&{d2},{S_STATUS},"OK"'
    e_by = lambda d1, d2, t: f'SUMIFS({E_AMT},{E_DATE},">="&{d1},{E_DATE},"<="&{d2},{E_TYPE},"{t}")'
    for i in range(12):
        r = REP_M0 + i
        ws.row_dimensions[r].height = 19
        ws[f"C{r}"] = f"=DATE(cfg_Year,{i + 1},1)"
        ws[f"D{r}"] = f"=EOMONTH(C{r},0)"
        ws[f"B{r}"] = f'=TEXT(C{r},"mmm")'
        d1, d2 = f"$C{r}", f"$D{r}"
        vals = {
            "E": f"=SUMIFS({S_REV},{by_date(d1, d2)})",
            "F": f"=SUMIFS({S_QTY},{by_date(d1, d2)})",
            "G": f"=SUMIFS({S_COGS},{by_date(d1, d2)})",
            "H": f"=E{r}-G{r}",
            "I": f'=IF(E{r}=0,"",H{r}/E{r})',
            "J": f'=SUMIFS({W_VAL},{W_DATE},">="&{d1},{W_DATE},"<="&{d2})',
            "K": "=" + e_by(d1, d2, "Fixed"),
            "L": "=" + e_by(d1, d2, "Labor"),
            "M": "=" + e_by(d1, d2, "Event Fees"),
            "N": "=" + e_by(d1, d2, "Variable"),
            "O": f"=SUM(K{r}:N{r})",
            "P": f"=H{r}-J{r}-O{r}",
            "Q": f'=IF(E{r}=0,"",P{r}/E{r})',
            "R": "=" + e_by(d1, d2, "Excluded - Stock Purchase"),
            "S": f"=SUMIFS({S_ING},{by_date(d1, d2)})",
        }
        style_calc(ws[f"B{r}"], bold=True)
        for c, fx in vals.items():
            style_calc(ws[f"{c}{r}"], PCT if c in "IQ" else (INT if c == "F" else MONEY), bold=(c == "P"))
            ws[f"{c}{r}"] = fx
    r = REP_TOT
    ws.row_dimensions[r].height = 22
    ws[f"B{r}"] = '="Total "&cfg_Year'
    style_header(ws[f"B{r}"])
    for c in "EFGHJKLMNOPRS":
        ws[f"{c}{r}"] = f"=SUM({c}{REP_M0}:{c}{REP_M1})"
        style_calc(ws[f"{c}{r}"], INT if c == "F" else MONEY, bold=True)
        ws[f"{c}{r}"].fill = fill(TEAL_LIGHT)
    for c, num in {"I": "H", "Q": "P"}.items():
        ws[f"{c}{r}"] = f'=IF(E{r}=0,"",{num}{r}/E{r})'
        style_calc(ws[f"{c}{r}"], PCT, bold=True)
        ws[f"{c}{r}"].fill = fill(TEAL_LIGHT)
    ws["B21"] = '=IF(E20=0,"No sales recorded for "&cfg_Year&" yet - add rows on the Sales Log.","")'
    style_label(ws["B21"], size=9, color=WARN_FONT)
    money_fmt(ws, [f"E{REP_M0}:E{REP_TOT}", f"G{REP_M0}:H{REP_TOT}", f"J{REP_M0}:P{REP_TOT}", f"R{REP_M0}:S{REP_TOT}"])

    # Selected period summary -------------------------------------------------
    group_band(ws, 23, "B", "G", "SELECTED PERIOD SUMMARY  (feeds the Dashboard)", INK)
    ws["B24"] = "=cfg_PeriodLabel"
    style_label(ws["B24"], bold=True, color=TEAL)
    p1, p2 = "cfg_Start", "cfg_End"
    rows = [
        ("rev", "Revenue", f"=SUMIFS({S_REV},{by_date(p1, p2)})", MONEY),
        ("items", "Items sold", f"=SUMIFS({S_QTY},{by_date(p1, p2)})", INT),
        ("cogs", "Cost of goods sold (COGS)", f"=SUMIFS({S_COGS},{by_date(p1, p2)})", MONEY),
        ("gp", "Gross profit", "={rev}-{cogs}", MONEY),
        ("gm", "Gross margin", '=IF({rev}=0,"",{gp}/{rev})', PCT),
        ("waste", "Waste", f'=SUMIFS({W_VAL},{W_DATE},">="&{p1},{W_DATE},"<="&{p2})', MONEY),
        ("fixed", "Fixed costs", "=" + e_by(p1, p2, "Fixed"), MONEY),
        ("labor", "Labour", "=" + e_by(p1, p2, "Labor"), MONEY),
        ("event", "Event fees", "=" + e_by(p1, p2, "Event Fees"), MONEY),
        ("var", "Other variable costs", "=" + e_by(p1, p2, "Variable"), MONEY),
        ("opex", "Total operating expenses", "={fixed}+{labor}+{event}+{var}", MONEY),
        ("net", "NET PROFIT (before tax)", "={gp}-{waste}-{opex}", MONEY),
        ("nm", "Net margin", '=IF({rev}=0,"",{net}/{rev})', PCT),
        ("lines", "Sales entries", f"=COUNTIFS({by_date(p1, p2)})", INT),
        ("avgp", "Average price per item", '=IF({items}=0,"",{rev}/{items})', MONEY),
        ("active", "Months with activity", f'=SUMPRODUCT((($E${REP_M0}:$E${REP_M1}>0)+($O${REP_M0}:$O${REP_M1}>0)>0)*'
                                           f'($C${REP_M0}:$C${REP_M1}>={p1})*($C${REP_M0}:$C${REP_M1}<={p2}))', INT),
        ("avgopex", "Average operating expenses per month", '=IF({active}=0,{opex},{opex}/{active})', MONEY),
        ("berev", "Break-even revenue for period", f"=IF(ISNUMBER({q(BE)}!$H$27),{q(BE)}!$H$27*cfg_Months,\"\")", MONEY),
        ("beprog", "Break-even progress", '=IF(OR({berev}="",N({berev})=0),"",{rev}/{berev})', PCT),
        ("stock", "Stock purchases (info only)", "=" + e_by(p1, p2, "Excluded - Stock Purchase"), MONEY),
        ("unclassified", "Expense rows needing a fix", f'=SUMPRODUCT(({rng(EXP, "I", FIRST, EXP_LAST)}<>"")*({rng(EXP, "I", FIRST, EXP_LAST)}<>"OK"))', INT),
    ]
    for i, (key, label, _f, _fmt) in enumerate(rows):
        PS[key] = f"$G${25 + i}"
    for i, (key, label, fx, fmt) in enumerate(rows):
        r = 25 + i
        ws[f"B{r}"] = label
        style_label(ws[f"B{r}"], bold=key in ("gp", "net", "rev"))
        ws.merge_cells(f"B{r}:F{r}")
        style_calc(ws[f"G{r}"], fmt, bold=key in ("gp", "net", "rev"))
        ws[f"G{r}"] = fx.format(**{k: v for k, v in PS.items()})
        if fmt == MONEY:
            money_fmt(ws, [f"G{r}"])
    for key in PS:
        add_name(wb, f"ps_{key}", f"{q(REP)}!{PS[key]}")
    # Cost structure (chart source for the Dashboard)
    group_band(ws, 48, "B", "G", "WHERE THE MONEY GOES (selected period) - Dashboard chart source", INK)
    for i, (label, key) in enumerate([("Food & packaging", "cogs"), ("Labour", "labor"), ("Fixed costs", "fixed"),
                                      ("Other variable", "var"), ("Event fees", "event"), ("Waste", "waste"),
                                      ("Net profit", "net")]):
        r = 49 + i
        ws[f"B{r}"] = label
        style_label(ws[f"B{r}"])
        ws.merge_cells(f"B{r}:F{r}")
        style_calc(ws[f"G{r}"], MONEY0)
        ws[f"G{r}"] = f"={PS[key]}"
    money_fmt(ws, ["G49:G55"], whole=True)

    # Expense breakdown by category (period) -------------------------------
    group_band(ws, 23, "J", "O", "OPERATING EXPENSES BY CATEGORY (selected period)", INK)
    headers(ws, 24, {"J": "Category", "M": "Cost type", "O": "% of opex"}, 20)
    ws.merge_cells("J24:L24")
    ws["N24"] = sym_header("Amount")
    style_header(ws["N24"])
    for i in range(20):
        r = 25 + i
        src = 6 + i
        ws.merge_cells(f"J{r}:L{r}")
        ws[f"J{r}"] = f'=IF(OR({q(LST)}!$F${src}="",LEFT({q(LST)}!$G${src},8)="Excluded"),"",{q(LST)}!$F${src})'
        ws[f"M{r}"] = f'=IF($J{r}="","",{q(LST)}!$G${src})'
        ws[f"N{r}"] = f'=IF($J{r}="","",SUMIFS({E_AMT},{E_CAT},$J{r},{E_DATE},">="&cfg_Start,{E_DATE},"<="&cfg_End))'
        ws[f"O{r}"] = f'=IF(OR($J{r}="",{PS["opex"]}=0),"",$N{r}/{PS["opex"]})'
        for c, fmt in {"J": None, "M": None, "N": MONEY, "O": PCT}.items():
            style_calc(ws[f"{c}{r}"], fmt)
    money_fmt(ws, ["N25:N44"])
    ws.freeze_panes = "C8"
    ws["B6"] = '="Year "&cfg_Year'
    style_label(ws["B6"], bold=True, color=TEAL)
    print_setup(ws, "A1:S56")


# ---------------------------------------------------------------------------
# Break-Even
# ---------------------------------------------------------------------------
def build_breakeven(wb, ws, sample):
    sheet_banner(ws, "Break-Even Calculator", "How many items you must sell each month to cover your costs - and to "
                 "hit your profit goal.", 9, [START, DASH, REP])
    set_widths(ws, {"B": 34, "C": 14, "D": 3, "E": 36, "F": 14, "G": 14, "H": 14, "I": 3})
    group_band(ws, 6, "B", "C", "1. MONTHLY FIXED COSTS (budget)")
    headers(ws, 7, {"B": "Cost"}, 20)
    ws["C7"] = sym_header("Per month")
    style_header(ws["C7"])
    for i in range(12):
        r = 8 + i
        style_input(ws[f"B{r}"])
        style_input(ws[f"C{r}"], MONEY)
        if i < len(SD.BREAK_EVEN_BUDGET):
            ws[f"B{r}"] = SD.BREAK_EVEN_BUDGET[i][0]
            if sample:
                ws[f"C{r}"] = SD.BREAK_EVEN_BUDGET[i][1]
    ws["B20"] = "Total monthly fixed costs"
    style_label(ws["B20"], bold=True)
    ws["C20"] = "=SUM(C8:C19)"
    style_calc(ws["C20"], MONEY, bold=True)
    ws["B21"] = "Actual avg opex / month (Expenses)"
    style_label(ws["B21"], size=9, color=MUTED)
    ws["C21"] = (f'=IF(COUNTIF({q(REP)}!$O${REP_M0}:$O${REP_M1},">0")=0,"",'
                 f'{q(REP)}!$O${REP_TOT}/COUNTIF({q(REP)}!$O${REP_M0}:$O${REP_M1},">0"))')
    style_calc(ws["C21"], MONEY)
    info_box(ws, "B23:C26", "Include every monthly cost that does NOT change with each item sold: rent, loan, insurance, "
             "planned wages, typical fuel. Do NOT include ingredients or packaging - those are in the cost per item.")
    dv_number(ws, "C8:C19", 0, None, "Monthly amount.", "Fixed cost")

    group_band(ws, 6, "E", "H", "2. PER-ITEM ECONOMICS")
    headers(ws, 7, {"E": "Driver", "F": "From your data", "G": "Your override", "H": "Used"}, 20)
    year_rev = f"{q(REP)}!$E${REP_TOT}"
    year_items = f"{q(REP)}!$F${REP_TOT}"
    year_cogs = f"{q(REP)}!$G${REP_TOT}"
    menu_o = rng(MENU, "O", FIRST, MENU_LAST)
    msum = lambda c: f'SUMIFS({rng(MENU, c, FIRST, MENU_LAST)},{menu_o},">0")'
    menu_cogs_avg = f'IFERROR(({msum("F")}+{msum("G")}+{msum("J")})/COUNTIF({menu_o},">0"),"")'
    drivers = [
        (8, "Average selling price per item", f'=IF({year_items}>0,{year_rev}/{year_items},IFERROR(AVERAGEIF({menu_o},">0"),""))', MONEY),
        (9, "Average COGS per item (food + packaging)", f'=IF({year_items}>0,{year_cogs}/{year_items},{menu_cogs_avg})', MONEY),
        (10, "Other variable costs (% of sales)", '=""', PCT),
        (11, "Service days per month", '=""', "0"),
        (12, "Target monthly profit", '=""', MONEY),
    ]
    for r, label, fx, fmt in drivers:
        ws[f"E{r}"] = label
        style_label(ws[f"E{r}"])
        style_calc(ws[f"F{r}"], fmt)
        ws[f"F{r}"] = fx if fx != '=""' else "-"
        if fx == '=""':
            ws[f"F{r}"].alignment = Alignment(horizontal="center")
        style_input(ws[f"G{r}"], fmt)
        style_calc(ws[f"H{r}"], fmt, bold=True)
        ws[f"H{r}"] = f'=IF(ISNUMBER(G{r}),G{r},IF(ISNUMBER(F{r}),F{r},0))'
    if sample:
        ws["G10"], ws["G11"], ws["G12"] = SD.OTHER_VARIABLE_PCT, SD.SERVICE_DAYS, SD.TARGET_PROFIT
    ws["E13"] = "Data source: Sales Log for the reporting year; if empty, the average of your menu prices/costs."
    style_label(ws["E13"], size=8, color=MUTED)
    dv_number(ws, "G8:G9", 0, None, "Leave blank to use your data.", "Override")
    dv_number(ws, "G10", 0, 0.9, "Card fees, delivery-app or event commissions, as % of sales.", "Variable %")
    dv_number(ws, "G11", 1, 31, "Days you trade per month.", "Service days", kind="whole")
    dv_number(ws, "G12", 0, None, "Profit you want after all costs.", "Target profit")

    group_band(ws, 16, "E", "H", "3. RESULTS", AMBER)
    res = [
        (17, "Contribution per item", '=H8-H9-H8*H10', MONEY),
        (18, "Contribution margin ratio", '=IF(H8<=0,"",H17/H8)', PCT),
        (19, "Monthly fixed costs", '=C20', MONEY),
        (20, "Break-even items per month", '=IF(C20<=0,"",IF(H17<=0,"Not reachable",ROUNDUP(C20/H17,0)))', INT),
        (21, "Break-even items per service day", '=IF(OR(NOT(ISNUMBER(H20)),H11<=0),"",ROUNDUP(H20/H11,0))', INT),
        (22, "Items needed for target profit", '=IF(OR(C20+H12<=0,H17<=0),"",ROUNDUP((C20+H12)/H17,0))', INT),
        (23, "Revenue needed for target profit", '=IF(NOT(ISNUMBER(H22)),"",(C20+H12)/H18)', MONEY),
        (24, "Actual average items sold per month", f'=IF(COUNTIF({q(REP)}!$F${REP_M0}:$F${REP_M1},">0")=0,"",{year_items}/COUNTIF({q(REP)}!$F${REP_M0}:$F${REP_M1},">0"))', INT),
        (25, "Margin of safety (actual vs break-even)", '=IF(OR(NOT(ISNUMBER(H24)),NOT(ISNUMBER(H20)),N(H24)=0),"",(H24-H20)/H24)', PCT),
        (26, "Status", '=IF(C20<=0,"Enter your monthly fixed costs (section 1)",IF(H8<=0,"Add menu prices or sales first",'
                       'IF(H17<=0,"Each sale loses money - raise prices or cut item costs",IF(NOT(ISNUMBER(H25)),"Break-even calculated - add sales to compare",'
                       'IF(H25>=0,"Above break-even by "&TEXT(H25,"0%"),"Below break-even by "&TEXT(-H25,"0%"))))))', None),
        (27, "Break-even revenue per month", '=IF(OR(C20<=0,N(H18)<=0),"",C20/H18)', MONEY),
    ]
    for r, label, fx, fmt in res:
        ws[f"E{r}"] = label
        style_label(ws[f"E{r}"], bold=r in (20, 27))
        ws.merge_cells(f"F{r}:G{r}") if False else None
        style_calc(ws[f"H{r}"], fmt, bold=r in (20, 27))
        ws[f"H{r}"] = fx
        ws.row_dimensions[r].height = 19
    ws["H26"].alignment = Alignment(horizontal="right", vertical="center")
    expr_rule(ws, "H26", 'ISNUMBER(SEARCH("Above",H26))', OK_FONT, OK_FILL, True)
    expr_rule(ws, "H26", 'OR(ISNUMBER(SEARCH("Below",H26)),ISNUMBER(SEARCH("loses",H26)))', WARN_FONT, WARN_FILL, True)
    for r in (20, 27):
        ws[f"H{r}"].fill = fill(TEAL_LIGHT)
        ws[f"H{r}"].font = font(11, True, TEAL_DARK)
    add_name(wb, "be_Revenue", f"{q(BE)}!$H$27")
    add_name(wb, "be_Items", f"{q(BE)}!$H$20")

    # Chart data -------------------------------------------------------------
    group_band(ws, 29, "B", "H", "4. BREAK-EVEN CHART", INK)
    group_band(ws, 50, "B", "F", "CHART DATA (automatic)", "6B7280")
    headers(ws, 51, {"B": "Items sold per month"}, 30)
    for c, t in {"C": "Revenue", "E": "Total costs", "F": "Fixed costs"}.items():
        style_header(ws[f"{c}51"])
        ws[f"{c}51"] = t
    for i in range(11):
        r = 52 + i
        ws[f"B{r}"] = f'=IF(ISNUMBER($H$20),ROUND($H$20*2*{i}/10,0),{i * 100})'
        ws[f"C{r}"] = f"=B{r}*$H$8"
        ws[f"E{r}"] = f"=$C$20+B{r}*($H$9+$H$8*$H$10)"
        ws[f"F{r}"] = "=$C$20"
        for c in "BCEF":
            style_calc(ws[f"{c}{r}"], INT if c == "B" else MONEY0)
    ws.column_dimensions["D"].width = 3
    # chart uses B (x), C, E, F
    chart = LineChart()
    chart.y_axis.title = "Per month"
    chart.x_axis.title = "Items sold per month"
    chart.height, chart.width = 9.5, 22
    for col_letter, color in (("C", "0F766E"), ("E", "C2410C"), ("F", "94A3B8")):
        ref = Reference(ws, range_string=f"'{BE}'!{col_letter}51:{col_letter}62")
        chart.add_data(ref, titles_from_data=True)
        s = chart.series[-1]
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = 28000
        s.smooth = False
    chart.set_categories(Reference(ws, range_string=f"'{BE}'!B52:B62"))
    chart.legend.position = "b"
    ws.add_chart(chart, "B30")
    money_fmt(ws, ["C8:C21", "F8:H9", "H12", "G12", "H17", "H19", "H23", "H27"])
    money_fmt(ws, ["C52:F62"], whole=True)
    print_setup(ws, "A1:I63")


# ---------------------------------------------------------------------------
# Event Predictor
# ---------------------------------------------------------------------------
def build_event(wb, ws, sample):
    sheet_banner(ws, "Event Profit Predictor", "Should you book this event? Forecast worst / base / best case, see the break-even "
                 "order count, then compare with what actually happened.", 8, [START, DASH, SALES])
    set_widths(ws, {"B": 44, "C": 15, "D": 15, "E": 15, "F": 3, "G": 46, "H": 3})
    group_band(ws, 6, "B", "E", "EVENT DETAILS")
    ev = SD.EVENT
    det = [(7, "Event name (= Location in Sales Log)", "name", None),
           (8, "Event date", "date", DATE), (9, "Notes", "notes", None)]
    for r, label, key, fmt in det:
        ws[f"B{r}"] = label
        style_label(ws[f"B{r}"])
        ws.merge_cells(f"C{r}:E{r}")
        style_input(ws[f"C{r}"], fmt)
        for c in "DE":
            style_input(ws[f"{c}{r}"])
        if sample:
            ws[f"C{r}"] = ev[key]
    dv_date(ws, "C8")

    group_band(ws, 11, "B", "E", "FORECAST - DEMAND (per scenario)", INK)
    headers(ws, 12, {"B": "Assumption", "C": "Worst case", "D": "Base case", "E": "Best case"}, 20)
    demand = [(13, "Expected attendance (people at the event)", "attendance", INT),
              (14, "% of attendees who buy from you", "buy_rate", PCT),
              (15, "Orders per buying customer", "orders_per_customer", "0.00"),
              (16, "Average order value", "aov", MONEY)]
    for r, label, key, fmt in demand:
        ws[f"B{r}"] = label
        style_label(ws[f"B{r}"])
        for j, c in enumerate("CDE"):
            style_input(ws[f"{c}{r}"], fmt, "right")
            if sample:
                ws[f"{c}{r}"] = ev[key][j]
    ws["B16"] = '="Average order value ("&cfg_Symbol&")"'
    dv_number(ws, "C13:E13", 0, None, "People expected at the event.", "Attendance", kind="whole")
    dv_number(ws, "C14:E14", 0, 1, "Share of attendees who buy, e.g. 5%.", "Buy rate")
    dv_number(ws, "C15:E15", 0, 20, "Usually 1.0", "Orders per customer")
    dv_number(ws, "C16:E16", 0, None, "Typical spend per order.", "Order value")

    group_band(ws, 18, "B", "E", "FORECAST - EVENT COSTS (same in every scenario)", INK)
    costs = [(19, "Event participation / pitch fee", "fee", MONEY),
             (20, "Organiser commission (% of sales)", "commission", PCT),
             (21, "Card processing (% of sales)", "card", PCT),
             (22, "Food & packaging cost (% of sales)", "food_pct", PCT),
             (23, "Staff working", "staff", "0"),
             (24, "Hours per staff member (incl. setup & travel)", "hours", "0.0"),
             (25, "Hourly wage incl. payroll taxes", "wage", MONEY),
             (26, "Transportation (mileage, tolls, parking)", "transport", MONEY),
             (27, "Fuel & generator", "fuel", MONEY),
             (28, "Other expenses (ice, extra permits, supplies)", "other", MONEY)]
    for r, label, key, fmt in costs:
        ws[f"B{r}"] = label
        style_label(ws[f"B{r}"])
        style_input(ws[f"C{r}"], fmt, "right")
        if sample:
            ws[f"C{r}"] = ev[key]
        ws.merge_cells(f"D{r}:E{r}")
    mo = rng(MENU, "O", FIRST, MENU_LAST)
    menu_food = (f'=IFERROR("Your menu average: "&TEXT((SUMIFS({rng(MENU, "F", FIRST, MENU_LAST)},{mo},">0")+SUMIFS({rng(MENU, "G", FIRST, MENU_LAST)},{mo},">0"))'
                 f'/SUMIF({mo},">0"),"0%"),"")')
    ws["D22"] = menu_food
    style_label(ws["D22"], size=8, color=MUTED)
    ws["D25"] = '="Settings labour rate: "&cfg_Symbol&TEXT(cfg_LaborRate,"0.00")'
    style_label(ws["D25"], size=8, color=MUTED)
    dv_number(ws, "C19", 0, None)
    dv_number(ws, "C20:C22", 0, 0.95, "Enter as a percentage.", "Percent")
    dv_number(ws, "C23", 0, 100, kind="whole")
    dv_number(ws, "C24:C28", 0, None)

    group_band(ws, 30, "B", "E", "FORECAST RESULTS", AMBER)
    headers(ws, 31, {"B": "Forecast (not actual results)", "C": "Worst case", "D": "Base case", "E": "Best case"}, 20)
    var_pct = "($C$20+$C$21+$C$22)"
    fixed_ev = "(N($C$19)+N($C$23)*N($C$24)*N($C$25)+N($C$26)+N($C$27)+N($C$28))"
    out = [
        (32, "Expected buying customers", "=ROUND(N({c}13)*N({c}14),0)", INT),
        (33, "Expected orders", "=ROUND({c}32*N({c}15),0)", INT),
        (34, "Revenue", "={c}33*N({c}16)", MONEY),
        (35, "Food & packaging", "={c}34*N($C$22)", MONEY),
        (36, "Commission + card fees", "={c}34*(N($C$20)+N($C$21))", MONEY),
        (37, "Contribution (after variable costs)", "={c}34-{c}35-{c}36", MONEY),
        (38, "Staffing cost", "=N($C$23)*N($C$24)*N($C$25)", MONEY),
        (39, "Fixed event costs (fee, staff, travel, fuel, other)", f"={fixed_ev}", MONEY),
        (40, "ESTIMATED PROFIT", "={c}37-{c}39", MONEY),
        (41, "Profit margin", '=IF({c}34=0,"",{c}40/{c}34)', PCT),
        (42, "Break-even orders", f'=IF(N({{c}}16)*(1-{var_pct})<=0,"n/a",ROUNDUP({fixed_ev}/(N({{c}}16)*(1-{var_pct})),0))', INT),
        (43, "Break-even attendance", '=IF(OR(NOT(ISNUMBER({c}42)),N({c}14)*N({c}15)=0),"n/a",ROUNDUP({c}42/(N({c}14)*N({c}15)),0))', INT),
        (44, "Orders above / below break-even", '=IF(ISNUMBER({c}42),{c}33-{c}42,"")', '+#,##0;-#,##0;0'),
    ]
    for r, label, fx, fmt in out:
        ws[f"B{r}"] = label
        style_label(ws[f"B{r}"], bold=r == 40)
        for c in "CDE":
            style_calc(ws[f"{c}{r}"], fmt, bold=r == 40)
            ws[f"{c}{r}"] = fx.format(c=c)
        ws.row_dimensions[r].height = 18
    for c in "CDE":
        ws[f"{c}40"].fill = fill(TEAL_LIGHT)
        ws[f"{c}40"].font = font(11, True, TEAL_DARK)
    expr_rule(ws, "C44:E44", 'AND(ISNUMBER(C44),C44<0)', WARN_FONT, None, True)
    ws["B46"] = "Verdict"
    style_label(ws["B46"], bold=True)
    ws.merge_cells("C46:E46")
    ws["C46"] = ('=IF(N(D16)=0,"Enter attendance, buy rate and order value to see a verdict",IF(C40>0,"LOW RISK - profitable even in the worst case",'
                 'IF(D40>0,"MODERATE RISK - profitable in the base case, loss in the worst case",IF(E40>0,"HIGH RISK - only the best case makes money",'
                 '"NOT VIABLE at these assumptions - negotiate the fee or skip"))))')
    style_calc(ws["C46"], bold=True)
    ws.row_dimensions[46].height = 24
    text_rule(ws, "C46", "C46", "LOW RISK", OK_FONT, OK_FILL)
    text_rule(ws, "C46", "C46", "MODERATE", CAUTION_FONT, CAUTION_FILL)
    expr_rule(ws, "C46", 'OR(ISNUMBER(SEARCH("HIGH",C46)),ISNUMBER(SEARCH("NOT VIABLE",C46)))', WARN_FONT, WARN_FILL, True)

    group_band(ws, 48, "B", "E", "ACTUAL RESULTS (after the event)", "374151")
    headers(ws, 49, {"B": "Actual", "C": "From your logs", "D": "Your entry", "E": "Used"}, 20)
    loc_ok = f'{S_LOC},$C$7,{S_STATUS},"OK"'
    act = [
        (50, "Actual revenue", f'=IF($C$7="","",SUMIFS({S_REV},{loc_ok}))', MONEY),
        (51, "Actual items sold", f'=IF($C$7="","",SUMIFS({S_QTY},{loc_ok}))', INT),
        (52, "Actual food & packaging (COGS)", f'=IF($C$7="","",SUMIFS({S_COGS},{loc_ok}))', MONEY),
        (53, "Actual fee + commission + card fees", '="-"', MONEY),
        (54, "Actual staffing cost", '="-"', MONEY),
        (55, "Actual transport, fuel & other", '="-"', MONEY),
    ]
    for r, label, fx, fmt in act:
        ws[f"B{r}"] = label
        style_label(ws[f"B{r}"])
        style_calc(ws[f"C{r}"], fmt)
        ws[f"C{r}"] = fx if fx != '="-"' else "-"
        ws[f"C{r}"].alignment = Alignment(horizontal="right")
        style_input(ws[f"D{r}"], fmt, "right")
        style_calc(ws[f"E{r}"], fmt, bold=True)
        ws[f"E{r}"] = f'=IF(ISNUMBER(D{r}),D{r},IF(ISNUMBER(C{r}),C{r},0))'
    if sample:
        ws["D53"], ws["D54"], ws["D55"] = round(ev["actual_fee"] + 0.026 * 8820.5, 2), ev["actual_staff"], ev["actual_other"]
    summ = [
        (56, "ACTUAL PROFIT", '=IF(AND(E50=0,E53+E54+E55=0),"",E50-E52-E53-E54-E55)', MONEY),
        (57, "Base-case forecast profit", '=D40', MONEY),
        (58, "Variance (actual - base forecast)", '=IF(E56="","",E56-D40)', '+#,##0.00;[Red]-#,##0.00;0.00'),
        (59, "Revenue vs base forecast", '=IF(OR(E56="",N(D34)=0),"",E50/D34)', PCT),
    ]
    for r, label, fx, fmt in summ:
        ws[f"B{r}"] = label
        style_label(ws[f"B{r}"], bold=r == 56)
        style_calc(ws[f"E{r}"], fmt, bold=True)
        ws[f"E{r}"] = fx
    dv_number(ws, "D50:D55", 0, None, "Optional: overrides the value from your logs.", "Actual")

    info_box(ws, "G6:G16",
             "HOW TO USE\n\n1. Name the event and fill the cream cells.\n2. Worst / base / best change only the demand "
             "(attendance, buy rate, order value).\n3. Check the verdict and the break-even orders before you pay a fee.\n\n"
             "After the event: log sales on the Sales Log with Location = the exact event name. Actual revenue, items "
             "and COGS appear here automatically. Type the actual costs in 'Your entry'.")
    info_box(ws, "G18:G26",
             "FORECAST vs ACTUAL\n\nForecast numbers are estimates and never appear on the Dashboard.\n\nTo count the "
             "event's real costs in your monthly profit, enter them on the Expenses sheet (Event & Pitch Fees, Staff Wages, "
             "Fuel...). This sheet only compares - it does not add anything to your totals, so nothing is counted twice.")
    money_fmt(ws, ["C16:E16", "C19", "C25:C28", "C34:E40", "C50:E50", "C52:E56", "E57:E58"])
    print_setup(ws, "A1:H60")


# ---------------------------------------------------------------------------
# Scenario Simulator
# ---------------------------------------------------------------------------
PRESETS = [
    ("Custom (use my levers)", None),
    ("Ingredient inflation +10%", (0.10, 0, 0, 0, 0)),
    ("Menu price increase +5%", (0, 0, 0.05, 0, 0)),
    ("Price +5%, volume -3%", (0, -0.03, 0.05, 0, 0)),
    ("Busy season (+25% volume)", (0, 0.25, 0, 0.10, 0)),
    ("Slow season (-30% volume)", (0, -0.30, 0, -0.15, 0)),
    ("Wage increase +8%", (0, 0, 0, 0.08, 0)),
    ("Event fees +20%", (0, 0, 0, 0, 0.20)),
    ("Perfect storm", (0.12, -0.15, 0, 0.08, 0.20)),
]


def build_simulator(wb, ws, sample):
    sheet_banner(ws, "Scenario Simulator", "What happens to profit if costs, prices or sales change? Your real data is "
                 "never modified.", 13, [START, DASH, REP])
    set_widths(ws, {"B": 40, "C": 15, "D": 15, "E": 15, "F": 12, "G": 3, "H": 28, "I": 10, "J": 10, "K": 10, "L": 10, "M": 10})
    group_band(ws, 6, "B", "E", "1. BASELINE - an average month")
    headers(ws, 7, {"B": "Baseline item", "C": "From your data", "D": "Your override", "E": "Baseline used"}, 20)
    rep = lambda c: f"{q(REP)}!${c}${REP_TOT}"
    nmonths = "$C$17"
    base = [
        (8, "Items sold per month", f'=IF({nmonths}=0,"",{rep("F")}/{nmonths})', INT),
        (9, "Average price per item", f'=IF(N({rep("F")})=0,"",{rep("E")}/{rep("F")})', MONEY),
        (10, "Ingredient cost per item", f'=IF(N({rep("F")})=0,"",{rep("S")}/{rep("F")})', MONEY),
        (11, "Packaging & other variable cost per item", f'=IF(N({rep("F")})=0,"",({rep("G")}-{rep("S")})/{rep("F")})', MONEY),
        (12, "Staff wages per month", f'=IF({nmonths}=0,"",{rep("L")}/{nmonths})', MONEY),
        (13, "Event & pitch fees per month", f'=IF({nmonths}=0,"",{rep("M")}/{nmonths})', MONEY),
        (14, "Other operating costs per month (fixed + variable)", f'=IF({nmonths}=0,"",({rep("K")}+{rep("N")})/{nmonths})', MONEY),
        (15, "Waste per month", f'=IF({nmonths}=0,"",{rep("J")}/{nmonths})', MONEY),
    ]
    for r, label, fx, fmt in base:
        ws[f"B{r}"] = label
        style_label(ws[f"B{r}"])
        style_calc(ws[f"C{r}"], fmt)
        ws[f"C{r}"] = fx
        style_input(ws[f"D{r}"], fmt, "right")
        style_calc(ws[f"E{r}"], fmt, bold=True)
        ws[f"E{r}"] = f"=IF(ISNUMBER(D{r}),D{r},N(C{r}))"
    dv_number(ws, "D8:D15", 0, None, "Leave blank to use your data.", "Override")
    ws["B17"] = "Months of sales data used for the averages"
    style_label(ws["B17"], size=9, color=MUTED)
    ws["C17"] = f'=COUNTIF({q(REP)}!$E${REP_M0}:$E${REP_M1},">0")'
    style_calc(ws["C17"], "0")
    ws.merge_cells("D17:E17")
    ws["D17"] = '=IF(C17=0,"No sales yet - type a baseline in the override column","Reporting year "&cfg_Year)'
    style_label(ws["D17"], size=8, color=MUTED)

    group_band(ws, 19, "B", "E", "2. CHOOSE A SCENARIO", AMBER)
    ws["B20"] = "Scenario"
    style_label(ws["B20"], bold=True)
    ws.merge_cells("C20:E20")
    style_input(ws["C20"])
    for c in "DE":
        style_input(ws[f"{c}20"])
    ws["C20"] = SD.SCENARIO_PRESET if sample else PRESETS[0][0]
    ws["C20"].font = font(11, True, INPUT_FONT)
    ws.row_dimensions[20].height = 22
    dv_list(ws, "=$H$23:$H$31", "C20", "Pick a preset, or 'Custom' to use your own levers below.", "Scenario")
    headers(ws, 22, {"B": "Lever (% change)", "C": "My levers\n(Custom)", "D": "Preset value", "E": "Applied"}, 30)
    levers = ["Ingredient prices", "Sales volume (items sold)", "Menu prices", "Labour costs", "Event fees"]
    for i, lv in enumerate(levers):
        r = 23 + i
        ws[f"B{r}"] = lv
        style_label(ws[f"B{r}"])
        style_input(ws[f"C{r}"], '+0.0%;-0.0%;0.0%', "right")
        style_calc(ws[f"D{r}"], '+0.0%;-0.0%;0.0%')
        pcol = "IJKLM"[i]
        ws[f"D{r}"] = f'=IFERROR(INDEX(${pcol}$23:${pcol}$31,MATCH($C$20,$H$23:$H$31,0)),0)'
        style_calc(ws[f"E{r}"], '+0.0%;-0.0%;0.0%', bold=True)
        ws[f"E{r}"] = f'=IF(OR($C$20=$H$23,$C$20=""),N(C{r}),D{r})'
    dv_number(ws, "C23:C27", -0.9, 5, "e.g. 10% or -15%", "Change")
    # presets table
    group_band(ws, 21, "H", "M", "PRESETS (editable)", "6B7280")
    headers(ws, 22, {"H": "Preset", "I": "Ingredients", "J": "Volume", "K": "Menu price", "L": "Labour", "M": "Event fees"}, 30)
    for i, (name, vals) in enumerate(PRESETS):
        r = 23 + i
        ws[f"H{r}"] = name
        style_input(ws[f"H{r}"]) if i else style_calc(ws[f"H{r}"])
        for j, c in enumerate("IJKLM"):
            if vals is None:
                style_calc(ws[f"{c}{r}"])
                ws[f"{c}{r}"] = "-"
                ws[f"{c}{r}"].alignment = Alignment(horizontal="center")
            else:
                style_input(ws[f"{c}{r}"], '+0%;-0%;0%', "right")
                ws[f"{c}{r}"] = vals[j]

    group_band(ws, 33, "B", "F", "3. IMPACT (per month)", INK)
    headers(ws, 34, {"B": "Metric", "C": "Baseline", "D": "Scenario", "E": "Change", "F": "Change %"}, 20)
    lev = {k: f"$E${23 + i}" for i, k in enumerate(["ing", "vol", "price", "lab", "fee"])}
    rows = [
        (35, "Items sold", "=E8", f"=C35*(1+{lev['vol']})", INT),
        (36, "Average price per item", "=E9", f"=C36*(1+{lev['price']})", MONEY),
        (37, "Revenue", "=C35*C36", "=D35*D36", MONEY),
        (38, "Ingredient costs", "=C35*E10", f"=D35*E10*(1+{lev['ing']})", MONEY),
        (39, "Packaging & other variable", "=C35*E11", "=D35*E11", MONEY),
        (40, "Gross profit", "=C37-C38-C39", "=D37-D38-D39", MONEY),
        (41, "Gross margin", '=IF(C37=0,"",C40/C37)', '=IF(D37=0,"",D40/D37)', PCT),
        (42, "Staff wages", "=E12", f"=E12*(1+{lev['lab']})", MONEY),
        (43, "Event & pitch fees", "=E13", f"=E13*(1+{lev['fee']})", MONEY),
        (44, "Other operating costs", "=E14", "=E14", MONEY),
        (45, "Waste (moves with ingredient prices)", "=E15", f"=E15*(1+{lev['ing']})", MONEY),
        (46, "NET PROFIT", "=C40-SUM(C42:C45)", "=D40-SUM(D42:D45)", MONEY),
        (47, "Net margin", '=IF(C37=0,"",C46/C37)', '=IF(D37=0,"",D46/D37)', PCT),
        (48, "Break-even items per month", '=IF(C36-E10-E11<=0,"n/a",ROUNDUP(SUM(C42:C45)/(C36-E10-E11),0))',
         f'=IF(D36-E10*(1+{lev["ing"]})-E11<=0,"n/a",ROUNDUP(SUM(D42:D45)/(D36-E10*(1+{lev["ing"]})-E11),0))', INT),
    ]
    for r, label, fb, fs, fmt in rows:
        ws[f"B{r}"] = label
        bold = r in (37, 40, 46)
        style_label(ws[f"B{r}"], bold=bold)
        style_calc(ws[f"C{r}"], fmt, bold)
        ws[f"C{r}"] = fb
        style_calc(ws[f"D{r}"], fmt, bold)
        ws[f"D{r}"] = fs
        if fmt == PCT:
            style_calc(ws[f"E{r}"], '+0.0" pts";-0.0" pts";0.0" pts"')
            ws[f"E{r}"] = f'=IF(OR(C{r}="",D{r}=""),"",(D{r}-C{r})*100)'
            style_calc(ws[f"F{r}"])
        else:
            style_calc(ws[f"E{r}"], '+#,##0.00;-#,##0.00;0.00' if fmt == MONEY else '+#,##0;-#,##0;0', bold)
            ws[f"E{r}"] = f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r})),D{r}-C{r},"")'
            style_calc(ws[f"F{r}"], '+0.0%;-0.0%;0.0%')
            ws[f"F{r}"] = f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r}),N(C{r})<>0),(D{r}-C{r})/ABS(C{r}),"")'
        ws.row_dimensions[r].height = 18
    for c in "CD":
        ws[f"{c}46"].fill = fill(TEAL_LIGHT)
        ws[f"{c}46"].font = font(11, True, TEAL_DARK)
    for good in ("E37:F37", "E40:F40", "E46:F46"):
        a = good.split(":")[0]
        expr_rule(ws, good, f'AND(ISNUMBER({a}),{a}<0)', WARN_FONT, None, True)
        expr_rule(ws, good, f'AND(ISNUMBER({a}),{a}>0)', OK_FONT, None, True)
    for bad in ("E38:F39", "E42:F45"):
        a = bad.split(":")[0]
        expr_rule(ws, bad, f'AND(ISNUMBER({a}),{a}>0)', WARN_FONT, None, False)

    ws["B50"] = "Average price needed to keep your baseline net profit"
    style_label(ws["B50"], bold=True)
    ws["D50"] = (f'=IF(OR(D35<=0,C35<=0),"",(C46+SUM(D42:D45))/D35+E10*(1+{lev["ing"]})+E11)')
    style_calc(ws["D50"], MONEY, True)
    ws["D50"].fill = fill(TEAL_LIGHT)
    ws["E50"] = '=IF(OR(D50="",N(C36)=0),"",D50/C36-1)'
    style_calc(ws["E50"], '+0.0%;-0.0%;0.0%', True)
    ws.merge_cells("B51:F51")
    ws["B51"] = ('=IF(D50="","",IF(E50>0.0005,"To keep today\'s profit in this scenario, raise your average price by "&TEXT(E50,"0.0%")&" (to "&cfg_Symbol&TEXT(D50,"0.00")&").",'
                 'IF(E50<-0.0005,"This scenario is more profitable: you could lower prices by up to "&TEXT(-E50,"0.0%")&" and keep today\'s profit.","Your profit is unchanged in this scenario.")))')
    style_label(ws["B51"], size=9, color=INFO_FONT, wrap=True)
    ws.row_dimensions[51].height = 30
    info_box(ws, "H34:M44",
             "HOW IT WORKS\n\nBaseline = your average month in the reporting year (or your overrides).\n\n"
             "The scenario multiplies the baseline by the applied levers. Price changes do not automatically change "
             "volume - use 'Price +5%, volume -3%' or your own levers to model customer reaction.\n\n"
             "Nothing on your other sheets changes.")
    money_fmt(ws, ["C9:E15", "C36:D40", "C42:D46", "D50"])
    # chart
    ws["H47"], ws["I47"], ws["J47"] = "Chart", "Baseline", "Scenario"
    for i, (lbl, r) in enumerate((("Revenue", 37), ("Gross profit", 40), ("Net profit", 46))):
        ws[f"H{48 + i}"] = lbl
        ws[f"I{48 + i}"] = f"=C{r}"
        ws[f"J{48 + i}"] = f"=D{r}"
    for c in "HIJ":
        for r in range(47, 51):
            ws[f"{c}{r}"].font = font(8, color="9CA3AF")
            ws[f"{c}{r}"].number_format = MONEY0
    ch = BarChart()
    ch.type = "col"
    group_band(ws, 52, "B", "F", "BASELINE vs SCENARIO (per month)", INK)
    ch.height, ch.width = 7, 13
    ch.add_data(Reference(ws, min_col=9, max_col=10, min_row=47, max_row=50), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=8, min_row=48, max_row=50))
    ch.series[0].graphicalProperties.solidFill = "94A3B8"
    ch.series[1].graphicalProperties.solidFill = TEAL
    ch.legend.position = "b"
    ch.y_axis.majorGridlines = None
    ch.height, ch.width = 7, 17
    ws.add_chart(ch, "B53")
    print_setup(ws, "A1:M68")


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------
def kpi_card(ws, top, left_col, label, value_formula, caption_formula, accent=TEAL):
    from openpyxl.utils import column_index_from_string as ci
    c0 = ci(left_col)
    c2 = c0 + 2
    L, R = left_col, col(c2)
    for r in (top, top + 1, top + 2):
        ws.merge_cells(f"{L}{r}:{R}{r}")
        for cc in range(c0, c2 + 1):
            cell = ws.cell(row=r, column=cc)
            cell.fill = fill(WHITE)
            side_top = Side(style="thick", color=accent) if r == top else None
            side_bot = Side(style="thin", color=LINE) if r == top + 2 else None
            cell.border = Border(top=side_top, bottom=side_bot,
                                 left=Side(style="thin", color=LINE) if cc == c0 else None,
                                 right=Side(style="thin", color=LINE) if cc == c2 else None)
    lab = ws[f"{L}{top}"]
    lab.value = label
    lab.font = font(8, True, MUTED)
    lab.alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
    val = ws[f"{L}{top + 1}"]
    val.value = value_formula
    val.font = font(20, True, INK)
    val.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    cap = ws[f"{L}{top + 2}"]
    cap.value = caption_formula
    cap.font = font(8, False, MUTED)
    cap.alignment = Alignment(horizontal="left", vertical="top", indent=1, wrap_text=True)
    return val


def money_text(ref, decimals=0):
    fmt = "#,##0" if decimals == 0 else "#,##0.00"
    return f'IF(N({ref})<0,"-","")&cfg_Symbol&TEXT(ABS(N({ref})),"{fmt}")'


def build_dashboard(wb, ws, sample):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = INK
    ws.column_dimensions["A"].width = 2.5
    for c in "BCDFGHJKLNOP":
        ws.column_dimensions[c].width = 10.5
    for c in "EIMQ":
        ws.column_dimensions[c].width = 2.2
    ws.column_dimensions["R"].width = 2.5
    for r in range(1, 70):
        for cc in range(1, 19):
            ws.cell(row=r, column=cc).fill = fill(PAGE)
    for cc in range(1, 19):
        ws.cell(row=1, column=cc).fill = fill(TEAL)
        ws.cell(row=2, column=cc).fill = fill(INK)
        ws.cell(row=3, column=cc).fill = fill(INK)
    ws.row_dimensions[1].height = 6
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 22
    ws.merge_cells("B2:L2")
    ws["B2"] = "=cfg_Business"
    ws["B2"].font = font(18, True, WHITE)
    ws["B2"].alignment = Alignment(vertical="center")
    ws.merge_cells("B3:L3")
    ws["B3"] = '="Executive Dashboard  ·  "&cfg_PeriodLabel&"  ·  "&cfg_Currency'
    ws["B3"].font = font(10, False, "C9D3DD")
    ws.merge_cells("M2:P2")
    ws["M2"] = "NoviraNest"
    ws["M2"].font = font(14, True, "5EEAD4")
    ws["M2"].alignment = Alignment(horizontal="right", vertical="center")
    ws.merge_cells("M3:P3")
    ws["M3"] = "Food Truck Profit & Pricing PRO"
    ws["M3"].font = font(8, False, "C9D3DD")
    ws["M3"].alignment = Alignment(horizontal="right")
    links = [("B4", START), ("D4", SET), ("F4", MENU), ("H4", SALES), ("J4", EXP), ("L4", BE), ("N4", EVT), ("P4", SIM)]
    for addr, sh in links:
        internal_link(ws[addr], sh, f"› {sh}", color=TEAL, size=8)
    ws.row_dimensions[5].height = 8
    ws.merge_cells("B6:P6")
    ws["B6"] = (f'=IF(ps_lines=0,"No sales recorded for "&cfg_PeriodLabel&" yet. Start with Settings, then add your menu, '
                f'and log sales on the Sales Log - this dashboard fills in automatically.",'
                f'IF(ISNUMBER(SEARCH("DEMO",cfg_Business)),"SAMPLE DATA: fictional demo business for illustration only - not real results. '
                f'Use the TEMPLATE file for your own business.",""))')
    ws["B6"].font = font(9, True, CAUTION_FONT)
    ws["B6"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    expr_rule(ws, "B6:P6", 'LEN($B$6)>0', CAUTION_FONT, CAUTION_FILL, True)
    ws.row_dimensions[6].height = 26
    ws.row_dimensions[7].height = 8

    # KPI rows
    for r in (8, 12):
        ws.row_dimensions[r].height = 18
        ws.row_dimensions[r + 1].height = 34
        ws.row_dimensions[r + 2].height = 26
    ws.row_dimensions[11].height = 10
    kpi_card(ws, 8, "B", "TOTAL REVENUE", "=" + money_text("ps_rev"),
             '=TEXT(ps_items,"#,##0")&" items sold · avg "&cfg_Symbol&TEXT(N(ps_avgp),"0.00")&" per item"')
    kpi_card(ws, 8, "F", "GROSS PROFIT", "=" + money_text("ps_gp"),
             '=IF(ps_gm="","Gross margin: -","Gross margin: "&TEXT(ps_gm,"0.0%"))&" · after food & packaging"')
    v = kpi_card(ws, 8, "J", "NET PROFIT", "=" + money_text("ps_net"),
                 '="After waste and all operating expenses (before tax)"', AMBER)
    expr_rule(ws, "J9", "ps_net<0", WARN_FONT, None, True)
    expr_rule(ws, "J9", "ps_net>0", "17663A", None, True)
    kpi_card(ws, 8, "N", "NET PROFIT MARGIN", '=IF(ps_nm="","-",TEXT(ps_nm,"0.0%"))',
             '="Net profit / revenue"', AMBER)
    expr_rule(ws, "N9", 'AND(ISNUMBER(ps_nm),ps_nm<0)', WARN_FONT, None, True)
    kpi_card(ws, 12, "B", "MONTHLY OPERATING EXPENSES", "=" + money_text("ps_avgopex"),
             '=IF(cfg_Months=1,"This month","Average of "&ps_active&" active month(s)")&" · total "&cfg_Symbol&TEXT(ps_opex,"#,##0")', INK)
    kpi_card(ws, 12, "F", "COST OF GOODS SOLD", "=" + money_text("ps_cogs"),
             '=IF(N(ps_rev)=0,"Food + packaging + other variable","= "&TEXT(ps_cogs/ps_rev,"0.0%")&" of revenue (food + packaging)")', INK)
    kpi_card(ws, 12, "J", "WASTE", "=" + money_text("ps_waste"),
             '=IF(N(ps_rev)=0,"Logged on the Waste Log",TEXT(ps_waste/ps_rev,"0.0%")&" of revenue · from the Waste Log")', INK)
    kpi_card(ws, 12, "N", "BREAK-EVEN PROGRESS", '=IF(ps_beprog="","Set up",TEXT(ps_beprog,"0%"))',
             '=IF(ps_beprog="","Enter fixed costs on Break-Even",REPT("█",MIN(10,ROUND(ps_beprog*10,0)))&REPT("░",MAX(0,10-ROUND(ps_beprog*10,0)))'
             '&"  of "&cfg_Symbol&TEXT(ps_berev,"#,##0"))', INK)
    expr_rule(ws, "N13", 'AND(ISNUMBER(ps_beprog),ps_beprog>=1)', "17663A", None, True)
    expr_rule(ws, "N13", 'AND(ISNUMBER(ps_beprog),ps_beprog<1)', WARN_FONT, None, True)
    ws.row_dimensions[15].height = 10

    # Section headers & charts
    def section(addr_range, text):
        ws.merge_cells(addr_range)
        a = addr_range.split(":")[0]
        ws[a] = text
        ws[a].font = font(10, True, INK)
        ws[a].alignment = Alignment(vertical="center")
        ws[a].border = Border(bottom=Side(style="medium", color=TEAL))
        for cell in ws[addr_range][0]:
            cell.border = Border(bottom=Side(style="medium", color=TEAL))

    ws.row_dimensions[16].height = 22
    section("B16:H16", '="MONTHLY PERFORMANCE  ·  "&cfg_Year')
    ws["B16"] = '="MONTHLY PERFORMANCE  ·  "&cfg_Year'
    section("J16:P16", '="WHERE THE MONEY GOES  ·  "&cfg_PeriodLabel')
    ws["J16"] = '="WHERE THE MONEY GOES  ·  "&cfg_PeriodLabel'
    bar = BarChart()
    bar.type = "col"
    bar.grouping = "clustered"
    bar.title = None
    bar.height, bar.width = 9.4, 13.2
    from openpyxl.chart.series import SeriesLabel
    for c, color, title in (("E", "0F766E", "Revenue"), ("H", "5EAAA8", "Gross profit"), ("P", "C2770E", "Net profit")):
        bar.add_data(Reference(wb[REP], range_string=f"'{REP}'!{c}8:{c}19"), titles_from_data=False)
        bar.series[-1].tx = SeriesLabel(v=title)
        bar.series[-1].graphicalProperties.solidFill = color
        bar.series[-1].graphicalProperties.line.noFill = True
    bar.set_categories(Reference(wb[REP], range_string=f"'{REP}'!B8:B19"))
    bar.legend.position = "b"
    bar.y_axis.majorGridlines.spPr = None
    bar.gapWidth = 60
    bar.y_axis.number_format = '#,##0'
    bar.y_axis.delete = False
    bar.x_axis.delete = False
    ws.add_chart(bar, "B17")
    hb = BarChart()
    hb.type = "bar"
    hb.title = None
    hb.height, hb.width = 9.4, 13.2
    hb.add_data(Reference(wb[REP], range_string=f"'{REP}'!G49:G55"), titles_from_data=False)
    hb.series[0].tx = SeriesLabel(v="Amount")
    hb.set_categories(Reference(wb[REP], range_string=f"'{REP}'!B49:B55"))
    hb.series[0].graphicalProperties.solidFill = INK
    hb.dataLabels = DataLabelList()
    hb.dataLabels.showVal = True
    hb.dataLabels.numFmt = '#,##0'
    hb.dataLabels.showCatName = False
    hb.dataLabels.showSerName = False
    hb.dataLabels.showLegendKey = False
    hb.dataLabels.position = "outEnd"
    hb.y_axis.majorGridlines = None
    hb.legend = None
    hb.y_axis.number_format = '#,##0'
    hb.x_axis.scaling.orientation = "maxMin"
    hb.gapWidth = 40
    hb.y_axis.delete = False
    hb.x_axis.delete = False
    ws.add_chart(hb, "J17")
    for r in range(17, 32):
        ws.row_dimensions[r].height = 18

    # Top tables
    ws.row_dimensions[32].height = 10
    ws.row_dimensions[33].height = 22
    section("B33:H33", '="BEST SELLERS  ·  "&cfg_PeriodLabel')
    section("J33:P33", '="MOST PROFITABLE ITEMS  ·  "&cfg_PeriodLabel')
    for start, cols, keycol, metric_hdr in (("B", ("B", "F", "G", "H"), "Z", "Items sold"),
                                            ("J", ("J", "N", "O", "P"), "AA", "Gross profit")):
        name_c, m_c, s_c, t_c = cols
        r = 34
        ws.merge_cells(f"{name_c}{r}:{col(ci_(m_c) - 1)}{r}")
        ws[f"{name_c}{r}"] = "#  Menu item"
        ws[f"{m_c}{r}"] = metric_hdr if keycol == "Z" else '="Gross profit ("&cfg_Symbol&")"'
        ws.merge_cells(f"{m_c}{r}:{col(ci_(m_c) + 1)}{r}")
        ws[f"{t_c}{r}"] = "Status" if keycol == "AA" else '="Revenue ("&cfg_Symbol&")"'
        if keycol == "Z":
            ws[f"{t_c}{r}"] = '="Revenue ("&cfg_Symbol&")"'
        for cc in range(ci_(name_c), ci_(t_c) + 1):
            cell = ws.cell(row=r, column=cc)
            cell.font = font(8, True, MUTED)
            cell.fill = fill(WHITE)
            cell.border = Border(bottom=Side(style="thin", color=LINE))
        key_rng = rng(MENU, keycol, FIRST, MENU_LAST)
        for k in range(1, 6):
            rr = 34 + k
            ws.row_dimensions[rr].height = 20
            idx = f"MATCH(LARGE({key_rng},{k}),{key_rng},0)"
            ok = f"LARGE({key_rng},{k})>=1" if keycol == "Z" else f"LARGE({key_rng},{k})>-1E+14"
            ok = f"AND({ok},INDEX({rng(MENU, 'V', FIRST, MENU_LAST)},{idx})>0)"
            ws.merge_cells(f"{name_c}{rr}:{col(ci_(m_c) - 1)}{rr}")
            ws[f"{name_c}{rr}"] = f'=IFERROR(IF({ok},"{k}.  "&INDEX({M_NAME},{idx}),IF({k}=1,"No sales in this period yet","")),"")'
            ws.merge_cells(f"{m_c}{rr}:{col(ci_(m_c) + 1)}{rr}")
            if keycol == "Z":
                ws[f"{m_c}{rr}"] = f'=IFERROR(IF({ok},INDEX({rng(MENU, "V", FIRST, MENU_LAST)},{idx}),""),"")'
                ws[f"{t_c}{rr}"] = f'=IFERROR(IF({ok},{money_text("INDEX(" + rng(MENU, "W", FIRST, MENU_LAST) + "," + idx + ")")},""),"")'
            else:
                ws[f"{m_c}{rr}"] = f'=IFERROR(IF({ok},INDEX({rng(MENU, "X", FIRST, MENU_LAST)},{idx}),""),"")'
                ws[f"{t_c}{rr}"] = f'=IFERROR(IF({ok},INDEX({rng(MENU, "T", FIRST, MENU_LAST)},{idx}),""),"")'
            for cc in range(ci_(name_c), ci_(t_c) + 1):
                cell = ws.cell(row=rr, column=cc)
                cell.fill = fill(WHITE)
                cell.font = font(9, cc == ci_(name_c), INK)
                cell.border = Border(bottom=Side(style="hair", color=LINE))
                cell.alignment = Alignment(vertical="center", horizontal="right" if cc >= ci_(m_c) else "left",
                                           indent=1 if cc == ci_(name_c) else 0)
            ws[f"{m_c}{rr}"].number_format = INT if keycol == "Z" else MONEY0
        text_rule(ws, f"{t_c}35:{t_c}39", f"{t_c}35", "Healthy", OK_FONT, None, mode="eq")
        expr_rule(ws, f"{t_c}35:{t_c}39", f'OR({t_c}35="LOSS",{t_c}35="Below target",{t_c}35="High food cost")', WARN_FONT, None, True)
    money_fmt(ws, ["N35:O39"], whole=True)

    # Alerts
    ws.row_dimensions[40].height = 10
    ws.row_dimensions[41].height = 22
    section("B41:P41", "ALERTS & DATA CHECKS")
    alerts = [
        (f'=COUNTIF({rng(MENU, "T", FIRST, MENU_LAST)},"LOSS")', "menu item(s) priced BELOW cost", MENU),
        (f'=COUNTIF({rng(MENU, "T", FIRST, MENU_LAST)},"Below target")', "menu item(s) below your target margin", MENU),
        (f'=COUNTIF({rng(INV, "P", FIRST, ING_LAST)},"Reorder")', "ingredient(s) at or below reorder level", INV),
        (f'=COUNTIF({rng(INV, "P", FIRST, ING_LAST)},"Recount (negative)")', "ingredient(s) with negative stock - recount", INV),
        (f'=SUMPRODUCT(({rng(REC, "G", FIRST, REC_LAST)}<>"")*({rng(REC, "G", FIRST, REC_LAST)}<>"OK"))', "recipe line(s) with a unit or setup problem", REC),
        (f'=SUMPRODUCT(({S_STATUS}<>"")*({S_STATUS}<>"OK"))', "sales row(s) that need fixing (not counted)", SALES),
        ("=ps_unclassified", "expense row(s) that need fixing (not counted)", EXP),
        (f'=SUMPRODUCT(({rng(WASTE, "H", FIRST, WASTE_LAST)}<>"")*({rng(WASTE, "H", FIRST, WASTE_LAST)}<>"OK"))', "waste row(s) that need fixing (not counted)", WASTE),
    ]
    for i, (fx, text, sheet) in enumerate(alerts):
        r = 42 + (i % 4)
        c_num, c_txt, c_lnk = ("B", "C", "H") if i < 4 else ("J", "K", "P")
        ws.row_dimensions[r].height = 20
        ws[f"{c_num}{r}"] = fx
        ws[f"{c_num}{r}"].font = font(12, True, OK_FONT)
        ws[f"{c_num}{r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"{c_num}{r}"].fill = fill(WHITE)
        ws.merge_cells(f"{c_txt}{r}:{col(ci_(c_lnk) - 1)}{r}")
        ws[f"{c_txt}{r}"] = text
        ws[f"{c_txt}{r}"].font = font(9, False, INK)
        ws[f"{c_txt}{r}"].alignment = Alignment(vertical="center")
        for cc in range(ci_(c_txt), ci_(c_lnk) + 1):
            ws.cell(row=r, column=cc).fill = fill(WHITE)
        internal_link(ws[f"{c_lnk}{r}"], sheet, "Open ›", size=8)
        ws[f"{c_lnk}{r}"].fill = fill(WHITE)
        expr_rule(ws, f"{c_num}{r}", f"{c_num}{r}>0", WARN_FONT, WARN_FILL, True)
    ws.row_dimensions[46].height = 10
    ws.merge_cells("B47:P47")
    ws["B47"] = ('="Profit model: Revenue - COGS (recipe cost x items sold) = Gross profit; minus waste and operating expenses = Net profit. '
                 'Ingredient purchases and owner\'s draw are excluded so nothing is counted twice.  ·  NoviraNest v' + VERSION + '"')
    ws["B47"].font = font(8, False, MUTED, True)
    ws["B47"].alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[47].height = 26
    print_setup(ws, "A1:R48")
    ws.page_setup.orientation = "landscape"
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.selectLockedCells = False


def ci_(letter):
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(letter)


# ---------------------------------------------------------------------------
# Start Here
# ---------------------------------------------------------------------------
def build_start(wb, ws, sample):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = AMBER
    set_widths(ws, {"A": 2.5, "B": 4, "C": 26, "D": 62, "E": 16, "F": 2.5})
    for cc in range(1, 7):
        ws.cell(row=1, column=cc).fill = fill(TEAL)
        for r in (2, 3, 4):
            ws.cell(row=r, column=cc).fill = fill(INK)
    ws.row_dimensions[1].height = 6
    ws.row_dimensions[2].height = 34
    ws.row_dimensions[3].height = 20
    ws.row_dimensions[4].height = 14
    ws.merge_cells("B2:E2")
    ws["B2"] = "Food Truck Profit & Pricing System PRO"
    ws["B2"].font = font(20, True, WHITE)
    ws["B2"].alignment = Alignment(vertical="center")
    ws.merge_cells("B3:E3")
    ws["B3"] = f"by NoviraNest  ·  Version {VERSION}  ·  Works in Microsoft Excel and Google Sheets  ·  No macros"
    ws["B3"].font = font(10, False, "C9D3DD")
    r = 6
    ws.merge_cells(f"B{r}:E{r}")
    if sample:
        ws[f"B{r}"] = ("THIS IS THE SAMPLE FILE. It contains a fictional demo business so you can explore every feature. "
                       "All names and numbers are invented and are NOT real results. For your own business, open the TEMPLATE file.")
        ws[f"B{r}"].fill, ws[f"B{r}"].font = fill(CAUTION_FILL), font(10, True, CAUTION_FONT)
    else:
        ws[f"B{r}"] = ("THIS IS YOUR CLEAN TEMPLATE. Every log is empty and ready for your data. Want to see it filled in "
                       "first? Open the SAMPLE file.")
        ws[f"B{r}"].fill, ws[f"B{r}"].font = fill(OK_FILL), font(10, True, OK_FONT)
    ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    ws.row_dimensions[r].height = 36

    def h(rr, text):
        ws.merge_cells(f"B{rr}:E{rr}")
        ws[f"B{rr}"] = text
        ws[f"B{rr}"].font = font(13, True, INK)
        ws[f"B{rr}"].border = Border(bottom=Side(style="medium", color=TEAL))
        for cc in range(2, 6):
            ws.cell(row=rr, column=cc).border = Border(bottom=Side(style="medium", color=TEAL))
        ws.row_dimensions[rr].height = 24

    h(8, "Quick setup  (about 30 minutes, once)")
    steps = [
        ("Settings", SET, "Business name, currency (USD or GBP), reporting year, labour rate, target margin and food cost %."),
        ("Inventory", INV, "List your ingredients: the unit you count them in, pack size and pack price. Cost per unit is automatic."),
        ("Recipe Builder", REC, "First add menu item names on Menu Pricing, then here list what goes into ONE serving. Units convert automatically."),
        ("Menu Pricing", MENU, "Add packaging, prep minutes and your menu price. See the suggested price, margin, markup and status."),
        ("Break-Even", BE, "Type your monthly fixed costs once. The Dashboard then shows your break-even progress."),
        ("Daily / weekly", SALES, "Log sales on the Sales Log, costs on Expenses and thrown-away food on the Waste Log. Then open the Dashboard."),
    ]
    for i, (name, sheet, text) in enumerate(steps):
        rr = 9 + i
        ws[f"B{rr}"] = i + 1
        ws[f"B{rr}"].font = font(12, True, WHITE)
        ws[f"B{rr}"].fill = fill(TEAL)
        ws[f"B{rr}"].alignment = Alignment(horizontal="center", vertical="center")
        internal_link(ws[f"C{rr}"], sheet, name, size=10)
        ws[f"C{rr}"].alignment = Alignment(vertical="center", indent=1)
        ws.merge_cells(f"D{rr}:E{rr}")
        ws[f"D{rr}"] = text
        ws[f"D{rr}"].font = font(10, False, CALC_FONT)
        ws[f"D{rr}"].alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[rr].height = 32

    h(16, "What the colours mean")
    leg = [("input", "Cream cell, blue text", "Type here. These are the only cells you need to fill in."),
           ("calc", "Grey cell", "Automatic formula. Please don't type over it."),
           ("warn", "Red / pink", "Something needs your attention (loss-making item, unit mismatch, bad row)."),
           ("info", "Blue box", "Instructions and explanations.")]
    for i, (kind, name, text) in enumerate(leg):
        rr = 17 + i
        c = ws[f"C{rr}"]
        c.value = name
        if kind == "input":
            style_input(c)
            c.protection = c.protection.copy(locked=True)
        elif kind == "calc":
            style_calc(c)
        elif kind == "warn":
            c.fill, c.font, c.border = fill(WARN_FILL), font(10, True, WARN_FONT), BOX
        else:
            style_info(c, 10)
        ws.merge_cells(f"D{rr}:E{rr}")
        ws[f"D{rr}"] = text
        ws[f"D{rr}"].font = font(10)
        ws.row_dimensions[rr].height = 20

    h(22, "Sheet directory")
    sheets = [
        (DASH, "KPI cards, monthly chart, expense breakdown, best sellers, most profitable items, alerts."),
        (SET, "Your business settings. Everything reads from here."),
        (MENU, "Smart pricing: cost per serving, suggested price, margin vs markup, food cost %, status, sales by item."),
        (REC, "Ingredients per serving with automatic unit conversion and mismatch protection."),
        (INV, "Ingredient prices, estimated stock on hand, stock value and reorder alerts."),
        (WASTE, "Food thrown away, valued automatically."),
        (SALES, "Every sale: revenue, cost and gross profit calculated for you (3,000 rows)."),
        (EXP, "Every business cost with a category (1,500 rows)."),
        (REP, "Profit & loss by month and the selected-period summary."),
        (BE, "Break-even items and revenue, target profit, chart."),
        (EVT, "Forecast an event (worst / base / best), then compare with actual results."),
        (SIM, "What-if: ingredient inflation, price changes, busy/slow season, wage rises."),
        (LST, "Units, expense categories and dropdown lists you can extend."),
    ]
    for i, (sheet, text) in enumerate(sheets):
        rr = 23 + i
        internal_link(ws[f"C{rr}"], sheet, f"› {sheet}", size=10)
        ws.merge_cells(f"D{rr}:E{rr}")
        ws[f"D{rr}"] = text
        ws[f"D{rr}"].font = font(10)
        ws[f"D{rr}"].alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[rr].height = 19

    h(37, "How the numbers work")
    defs = [
        ("Revenue", "Items sold x price charged (menu price unless you typed a different price on the Sales Log)."),
        ("COGS", "Cost of goods sold = items sold x (ingredients + packaging + other variable cost per serving). Labour is NOT in COGS."),
        ("Gross profit", "Revenue - COGS.  Gross margin = gross profit / revenue."),
        ("Net profit", "Gross profit - waste - operating expenses (fixed, labour, event fees, other variable). Before income tax."),
        ("Not double counted", "Ingredient purchases and owner's draw can be logged on Expenses but are excluded from profit: "
                               "the food you used is already in COGS. Prep labour on Menu Pricing is a pricing allowance only; real wages go in Expenses."),
        ("Margin vs markup", "Margin = profit / PRICE.  Markup = profit / COST.  Cost 4.00, price 10.00 -> margin 60%, markup 150%."),
        ("Food cost %", "Ingredient cost / menu price."),
        ("Break-even", "Monthly fixed costs / contribution per item (price - item cost - % fees)."),
        ("Currency", "USD or GBP changes symbols only. Amounts are never converted with an exchange rate."),
        ("Forecasts", "Event Predictor and Scenario Simulator are estimates. They never change your actual figures."),
    ]
    for i, (term, text) in enumerate(defs):
        rr = 38 + i
        ws[f"C{rr}"] = term
        ws[f"C{rr}"].font = font(10, True, TEAL_DARK)
        ws[f"C{rr}"].alignment = Alignment(vertical="top", indent=1)
        ws.merge_cells(f"D{rr}:E{rr}")
        ws[f"D{rr}"] = text
        ws[f"D{rr}"].font = font(10)
        ws[f"D{rr}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[rr].height = 30 if len(text) > 80 else 18
    ws.row_dimensions[42].height = 44

    h(49, "Good to know")
    tips = [
        "Need more rows? Select the last row of a log, copy it, and paste it below - the formulas come with it. "
        "Then extend the ranges named in the Sales/Expense formulas, or simply use the 3,000 / 1,500 rows provided.",
        "Google Sheets: File > Import > Upload, or open the .xlsx from Google Drive and choose File > Save as Google Sheets. "
        "See the Quick-Start Guide PDF for what looks different in Google Sheets.",
        "Rows with a red 'Check' are NOT counted until fixed - the Dashboard 'Alerts' panel tells you how many.",
        "Back up your file regularly (Excel: Save a copy; Google Sheets: File > Version history).",
        "This tool provides management estimates, not tax or accounting advice.",
    ]
    for i, t in enumerate(tips):
        rr = 50 + i
        ws[f"B{rr}"] = "•"
        ws[f"B{rr}"].font = font(11, True, TEAL)
        ws[f"B{rr}"].alignment = Alignment(horizontal="center", vertical="top")
        ws.merge_cells(f"C{rr}:E{rr}")
        ws[f"C{rr}"] = t
        ws[f"C{rr}"].font = font(10)
        ws[f"C{rr}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[rr].height = 30
    ws.merge_cells("B56:E56")
    ws["B56"] = (f"Version {VERSION} ({RELEASE_DATE})  ·  © NoviraNest. Licensed for use by one business. "
                 "Do not resell or share the file. See LICENSE.txt.")
    ws["B56"].font = font(8, False, MUTED, True)
    print_setup(ws, "A1:F56", landscape=False)


# ---------------------------------------------------------------------------
PREVIEW_AREAS = {
    START: "A1:F56", DASH: "A1:R48", SET: "A1:F34", MENU: "A1:Y24", REC: "A1:L32", INV: "A1:Q40",
    WASTE: "A1:J30", SALES: "A1:N34", EXP: "A1:I34", REP: "A1:S46", BE: "A1:I47", EVT: "A1:H60",
    SIM: "A1:M68", LST: "A1:V41",
}


def build(kind, path, preview=False):
    sample = kind == "sample"
    data = SD.build_dataset() if sample else None
    wb = Workbook()
    order = [START, DASH, SET, MENU, REC, INV, WASTE, SALES, EXP, REP, BE, EVT, SIM, LST]
    sheets = {}
    first = wb.active
    first.title = order[0]
    sheets[order[0]] = first
    for name in order[1:]:
        sheets[name] = wb.create_sheet(name)
    build_lists(wb, sheets[LST])
    build_settings(wb, sheets[SET], sample)
    build_inventory(wb, sheets[INV], data)
    build_recipes(wb, sheets[REC], sample)
    build_menu(wb, sheets[MENU], sample)
    build_sales(wb, sheets[SALES], data)
    build_expenses(wb, sheets[EXP], data)
    build_waste(wb, sheets[WASTE], data)
    build_report(wb, sheets[REP])
    build_breakeven(wb, sheets[BE], sample)
    build_event(wb, sheets[EVT], sample)
    build_simulator(wb, sheets[SIM], sample)
    build_dashboard(wb, sheets[DASH], sample)
    build_start(wb, sheets[START], sample)
    for ws in wb.worksheets:
        place_nav(ws)
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 100
        ws.sheet_view.selection[0].activeCell = "A1"
        ws.sheet_view.selection[0].sqref = "A1"
    wb.properties.title = "Food Truck Profit & Pricing System PRO"
    wb.properties.creator = "NoviraNest"
    wb.properties.subject = f"Version {VERSION}"
    wb.properties.keywords = "food truck, pricing, profit, NoviraNest"
    wb.calculation.fullCalcOnLoad = True
    if preview:  # one page per sheet, used only for rendering screenshots
        for ws in wb.worksheets:
            ws.print_area = PREVIEW_AREAS[ws.title]
            ws.page_setup.fitToHeight = 1
            ws.page_setup.fitToWidth = 1
            ws.page_setup.orientation = "landscape" if ws.title not in (START, SET) else "portrait"
    wb.save(path)
    return path


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
    print("built", sys.argv[2])
