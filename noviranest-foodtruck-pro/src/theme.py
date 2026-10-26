"""NoviraNest visual identity and openpyxl style helpers."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.formatting.rule import Rule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

FONT = "Arial"

# Palette ---------------------------------------------------------------
INK = "1E2A38"        # deep slate - titles, header bands
TEAL = "0F766E"       # primary brand
TEAL_DARK = "0B5A54"
TEAL_LIGHT = "E3F1EF"
AMBER = "C2770E"      # accent
MUTED = "6B7280"      # secondary text
PAGE = "F3F5F7"       # dashboard canvas
WHITE = "FFFFFF"
LINE = "DDE2E7"

INPUT_FILL, INPUT_FONT, INPUT_LINE = "FFF6E0", "1D3F8F", "EAD9AE"
CALC_FILL, CALC_FONT = "F3F5F7", "1F2937"
WARN_FILL, WARN_FONT = "FDECEC", "B42318"
OK_FILL, OK_FONT = "E6F4EA", "17663A"
CAUTION_FILL, CAUTION_FONT = "FFF1D6", "8A4B00"
INFO_FILL, INFO_FONT = "EAF2FB", "1E4E8C"

# Number formats -------------------------------------------------------
MONEY = '#,##0.00;[Red]-#,##0.00;0.00'
MONEY0 = '#,##0;[Red]-#,##0;0'
PCT = '0.0%;[Red]-0.0%;0.0%'
QTY = '#,##0.###'
INT = '#,##0'
DATE = 'dd mmm yyyy'

CURRENCY_FORMATS = {
    "USD": ('"$"#,##0.00;[Red]-"$"#,##0.00;"$"0.00', '"$"#,##0;[Red]-"$"#,##0;"$"0'),
    "GBP": ('"£"#,##0.00;[Red]-"£"#,##0.00;"£"0.00', '"£"#,##0;[Red]-"£"#,##0;"£"0'),
}

_thin = Side(style="thin", color=LINE)
_input_side = Side(style="thin", color=INPUT_LINE)
BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
INPUT_BOX = Border(left=_input_side, right=_input_side, top=_input_side, bottom=_input_side)
BOTTOM = Border(bottom=Side(style="thin", color=LINE))


def font(size=10, bold=False, color=CALC_FONT, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def style_input(c, fmt=None, align=None):
    c.fill = fill(INPUT_FILL)
    c.font = font(10, color=INPUT_FONT)
    c.border = INPUT_BOX
    c.alignment = Alignment(horizontal=align, vertical="center")
    c.protection = c.protection.copy(locked=False)
    if fmt:
        c.number_format = fmt


def style_calc(c, fmt=None, bold=False, align=None):
    c.fill = fill(CALC_FILL)
    c.font = font(10, bold=bold)
    c.border = BOX
    c.alignment = Alignment(horizontal=align, vertical="center")
    if fmt:
        c.number_format = fmt


def style_header(c, dark=True, wrap=True):
    c.fill = fill(INK if dark else TEAL)
    c.font = font(9, bold=True, color=WHITE)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)
    c.border = Border(left=Side(style="thin", color="3A4757"), right=Side(style="thin", color="3A4757"))


def style_group(c, color=TEAL):
    c.fill = fill(color)
    c.font = font(9, bold=True, color=WHITE)
    c.alignment = Alignment(horizontal="center", vertical="center")


def style_info(c, size=9, wrap=True):
    c.fill = fill(INFO_FILL)
    c.font = font(size, color=INFO_FONT)
    c.alignment = Alignment(wrap_text=wrap, vertical="center")


def style_label(c, bold=False, size=10, color=CALC_FONT, wrap=False):
    c.font = font(size, bold=bold, color=color)
    c.alignment = Alignment(vertical="center", wrap_text=wrap)


def internal_link(c, sheet, text=None, color=TEAL, size=9, bold=True):
    if text is not None:
        c.value = text
    c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{sheet}'!A1", display=str(c.value))
    c.font = Font(name=FONT, size=size, bold=bold, color=color, underline="single")


def sheet_banner(ws, title, subtitle, last_col, nav=None, tab_color=TEAL):
    """Rows 1-4: brand strip, title, subtitle, navigation."""
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab_color
    ws.column_dimensions["A"].width = 2.5
    ws.row_dimensions[1].height = 6
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 8
    for col in range(1, last_col + 2):
        ws.cell(row=1, column=col).fill = fill(TEAL)
    t = ws.cell(row=2, column=2, value=title)
    t.font = font(18, bold=True, color=INK)
    t.alignment = Alignment(vertical="center")
    s = ws.cell(row=3, column=2, value=subtitle)
    s.font = font(9, color=MUTED, italic=True)
    brand = ws.cell(row=2, column=last_col, value="NoviraNest")
    brand.font = font(8, bold=True, color=TEAL)
    brand.alignment = Alignment(horizontal="right", vertical="center")
    ws._nav = nav or []


def place_nav(ws):
    """Place navigation links on row 4 so each label fits in the columns it spans."""
    col_i, budget = 2, 0.0
    for name in getattr(ws, "_nav", []):
        internal_link(ws.cell(row=4, column=col_i), name, f"› {name}")
        need = len(name) * 1.05 + 4
        width = 0.0
        while width < need:
            w = ws.column_dimensions[get_column_letter(col_i)].width or 8.43
            width += w
            col_i += 1


def set_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def currency_cf(ws, ranges, symbol_ref, whole=False):
    """Excel: show $ or £ on money cells depending on Settings currency.
    Google Sheets ignores number formats in conditional formatting; headers carry the symbol there."""
    for code, (fmt2, fmt0) in CURRENCY_FORMATS.items():
        fmt = fmt0 if whole else fmt2
        num_id = {('USD', False): 300, ('USD', True): 301, ('GBP', False): 302, ('GBP', True): 303}[(code, whole)]
        dxf = DifferentialStyle(numFmt=NumberFormat(numFmtId=num_id, formatCode=fmt))
        rule = Rule(type="expression", dxf=dxf, formula=[f'{symbol_ref}="{code}"'])
        for rng in ranges:
            ws.conditional_formatting.add(rng, rule)


def text_rule(ws, rng, first_cell, text, fg, bg, bold=True, mode="contains"):
    if mode == "contains":
        formula = f'ISNUMBER(SEARCH("{text}",{first_cell}))'
    else:
        formula = f'{first_cell}="{text}"'
    dxf = DifferentialStyle(font=Font(color=fg, bold=bold), fill=PatternFill(start_color=bg, end_color=bg, fill_type="solid"))
    ws.conditional_formatting.add(rng, Rule(type="expression", dxf=dxf, formula=[formula]))


def expr_rule(ws, rng, formula, fg=None, bg=None, bold=False):
    kw = {}
    if fg:
        kw["font"] = Font(color=fg, bold=bold)
    if bg:
        kw["fill"] = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
    ws.conditional_formatting.add(rng, Rule(type="expression", dxf=DifferentialStyle(**kw), formula=[formula]))


def col(n):
    return get_column_letter(n)
